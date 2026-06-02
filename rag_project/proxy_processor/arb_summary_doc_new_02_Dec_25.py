from proxy_processor.merger_background_8_cleaned_format import ProxyBackgroundAnalyzer, DOCXFormatter
import os
import sys
import json
import re
import time
import openai
import anthropic
import pinecone
from dotenv import load_dotenv
import logging
from typing import List, Dict, Any, Tuple
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Import merger background analyzer
sys.path.append(os.path.dirname(os.path.abspath(__file__)))

# Load environment variables
load_dotenv()

# Set up logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s',
    handlers=[
        logging.StreamHandler(sys.stdout),
        logging.FileHandler('query_debugging.log', mode='w', encoding='utf-8')
    ]
)
logger = logging.getLogger(__name__)


class QueryProcessor:
    """Service to process questions and search for relevant content in Pinecone"""

    def __init__(self):
        logger.info("Initializing QueryProcessor")
        # Initialize OpenAI client (for embeddings)
        self.openai_client = openai.OpenAI(
            api_key=os.environ.get("OPENAI_API_KEY_SEC_FILING")
        )
        logger.info(
            f"OpenAI API key set: {'Yes' if os.environ.get('OPENAI_API_KEY_SEC_FILING') else 'No'}"
        )

        # Initialize Claude client (for Q&A)
        self.claude_client = anthropic.Anthropic(
            api_key=os.environ.get("ANTHROPIC_API_KEY")
        )
        self.claude_model = "claude-sonnet-4-5-20250929"
        logger.info(
            f"Anthropic API key set: {'Yes' if os.environ.get('ANTHROPIC_API_KEY') else 'No'}"
        )
        logger.info(f"Using Claude model: {self.claude_model}")

        # Initialize Pinecone
        pc = pinecone.Pinecone(api_key=os.environ.get("PINECONE_API_KEY"))
        logger.info(
            f"Pinecone API key set: {'Yes' if os.environ.get('PINECONE_API_KEY') else 'No'}"
        )

        # Get index
        self.index_name = os.environ.get(
            "PINECONE_INDEX_NAME_PROXY", "contract-chunks")
        logger.info(f"Using Pinecone index: {self.index_name}")

        try:
            # Connect to the index
            self.index = pc.Index(self.index_name)
            logger.info(f"Connected to Pinecone index: {self.index_name}")

        except Exception as e:
            logger.error(f"Error with Pinecone initialization: {str(e)}")
            raise

    def create_query_embedding(self, query: str) -> List[float]:
        """Create an embedding for the query using OpenAI"""
        try:
            logger.info(f"Creating embedding for query: {query}")
            response = self.openai_client.embeddings.create(
                input=query,
                model="text-embedding-3-large"
            )
            logger.info("Query embedding created successfully")
            return response.data[0].embedding
        except Exception as e:
            logger.error(f"Error creating query embedding: {str(e)}")
            raise

    def search_preamble_chunks(self, query_embedding: List[float], deal_id: str, top_k: int = 8) -> List[Dict]:
        """Search for preamble chunks specifically"""
        try:
            logger.info(f"Searching for preamble chunks (Deal ID: {deal_id})")
            search_response = self.index.query(
                vector=query_embedding,
                top_k=top_k,
                include_metadata=True,
                filter={
                    "deal_id": deal_id,
                    "title": {"$eq": "Preamble"}
                }
            )

            results = []
            for match in search_response.matches:
                result = {
                    'score': match.score,
                    'text': match.metadata['original_text'],
                    'id': match.id
                }
                results.append(result)

            logger.info(f"Found {len(results)} preamble chunks")
            return results
        except Exception as e:
            logger.error(f"Error searching preamble chunks: {str(e)}")
            return []

    def search_general_chunks(self, query_embedding: List[float], deal_id: str, top_k: int = 8) -> List[Dict]:
        """Search for general chunks with only deal_id filter"""
        try:
            logger.info(f"Searching for general chunks (Deal ID: {deal_id})")
            search_response = self.index.query(
                vector=query_embedding,
                top_k=top_k,
                include_metadata=True,
                filter={
                    "deal_id": deal_id
                }
            )

            results = []
            for match in search_response.matches:
                result = {
                    'score': match.score,
                    'text': match.metadata['original_text'],
                    'id': match.id
                }
                results.append(result)

            logger.info(f"Found {len(results)} general chunks")
            return results
        except Exception as e:
            logger.error(f"Error searching general chunks: {str(e)}")
            return []

    def search_similar_chunks(self, query: str, deal_id: str, top_k: int = 8) -> List[Dict]:
        """Search for similar chunks using multiple search strategies"""
        try:
            # Create query embedding
            query_embedding = self.create_query_embedding(query)
            all_results = []

            # 1. Always search for preamble chunks
            logger.info("Step 1: Searching for preamble chunks")
            preamble_results = self.search_preamble_chunks(
                query_embedding, deal_id, top_k)
            all_results.extend(preamble_results)

            # 2. Always search for general chunks with deal_id filter
            logger.info("Step 2: Searching for general chunks")
            general_results = self.search_general_chunks(
                query_embedding, deal_id, top_k)
            all_results.extend(general_results)

            # 3. Remove duplicates based on chunk ID
            logger.info("Step 3: Removing duplicates")
            seen_ids = set()
            deduplicated_results = []
            for result in all_results:
                if result['id'] not in seen_ids:
                    seen_ids.add(result['id'])
                    deduplicated_results.append(result)

            logger.info(
                f"Total results before deduplication: {len(all_results)}")
            logger.info(f"Final unique results: {len(deduplicated_results)}")
            return deduplicated_results

        except Exception as e:
            logger.error(f"Error searching Pinecone: {str(e)}")
            raise

    def format_results(self, results: List[Dict]) -> str:
        """Format search results into a readable string"""
        if not results:
            return "No relevant information found."

        formatted_output = "Here are the most relevant sections:\n\n"
        for i, result in enumerate(results, 1):
            formatted_output += f"Match {i} (Score: {result['score']:.3f}):\n"
            formatted_output += f"{result['text']}\n\n"

        return formatted_output

    def get_claude_response(self, query: str, results: List[Dict]) -> str:
        """Get a refined answer from Claude based on the search results"""
        try:
            # Prepare context from results
            context = "\n\n".join([f"Content {i+1}:\n{r['text']}"
                                   for i, r in enumerate(results)])

            # Construct the prompt
            prompt = f"""You are a merger arbitrage analyst at a hedge fund reviewing an SEC proxy filing. Answer questions with the specificity a risk arb analyst needs: exact dates, dollar amounts, bid trajectories, and factors material to deal completion risk and timeline.
            
            Based on the following content from an SEC filing document, please answer the question.
If the answer cannot be fully determined from the provided content, please mention that.

Question: {query}

Relevant content from the document:
{context}

Please provide a clear, concise answer based on the above content. If there are any uncertainties or if the information seems incomplete, please note that in your response."""

            logger.info(f"Sending request to Claude ({self.claude_model})")
            response = self.claude_client.messages.create(
                model=self.claude_model,
                max_tokens=4000,
                messages=[{"role": "user", "content": prompt}],
                temperature=0,
                timeout=120
            )

            answer = response.content[0].text.strip()
            logger.info(f"Received response from Claude ({len(answer)} chars)")
            return answer

        except Exception as e:
            logger.error(f"Error getting Claude response: {str(e)}")
            raise

    def generate_arb_summary(self, comprehensive_answer: str, query: str) -> str:
        """Generate a concise arbitrage bullet-point summary from the answer."""
        print("\n📊 Generating arbitrage summary...")
        logger.info("Generating arbitrage summary")

        prompt = f"""You are a professional legal/financial analyst drafting a concise DOC report.

INPUT:
Question: {query}
Answer: {comprehensive_answer}

TASK:
Transform the input into a shortened DOC-style structured output with the following rules:
1. Begin with a bold **Answer Header** (use the question as the header), don't add **Q**, **A** like formating only given question as the header.
2. There should be one paragraph only with concise fact, obligation, or key item.
3. Keep the total response brief.
4. Maintain a professional, memo-like tone (neutral, factual).
5. Remove redundant or minor details — focus on essentials.

OUTPUT:
Generate a **bullet-point Q&A format** that is DOC-ready and suitable for an executive summary.
"""

        try:
            logger.info(f"Sending arbitrage summary request to Claude")
            response = self.claude_client.messages.create(
                model=self.claude_model,
                max_tokens=512,
                messages=[{"role": "user", "content": prompt}],
                temperature=0,
                timeout=120
            )
            summary = response.content[0].text.strip()

            # Remove markdown code blocks if present
            if summary.startswith("```"):
                summary = summary.split("```")[1]
                if summary.startswith("markdown") or summary.startswith("text"):
                    summary = "\n".join(summary.split("\n")[1:])
                summary = summary.strip()

            logger.info(f"Arbitrage summary generated ({len(summary)} chars)")
            print(f"✅ Summary generated\n")
            return summary

        except Exception as e:
            error_msg = str(e)
            logger.error(f"Error generating summary: {error_msg}")
            print(f"❌ Error: {error_msg}")
            return f"Error generating summary: {error_msg}"


def convert_txt_to_docx(txt_file: str, deal_id: str, merger_background_results: Dict[str, Any] = None) -> str:
    """Convert the text file(s) to a Word document with Q&A section and merger background analysis."""
    try:
        # Read the first text file (JSON questions)
        with open(txt_file, "r", encoding="utf-8") as f:
            content = f.read()

        # Create a new Word document
        doc = Document()
        formatter = DOCXFormatter(doc)

        # Add title with custom formatting
        formatter.add_title('Proxy Summary')
        # formatter.add_metadata("claude-sonnet-4-5-20250929")
        # doc.add_page_break()

        # Add spacing
        doc.add_paragraph()

        # Process first file (JSON questions)
        if content.strip():
            _process_content_section(doc, content)

        # Add merger background analysis if available
        if merger_background_results and merger_background_results.get('success'):
            # Add page break before merger background sections
            doc.add_page_break()

            # =====================================================================
            # CLIENT DELIVERABLES SECTION
            # =====================================================================
            doc.add_heading(
                'Merger Background Analysis - Client Deliverables', 1)

            # Add document purpose
            # purpose = ("This document contains the specific deliverables requested in your specifications: "
            #            "chronological summary (numbered format), complete bidder census, bid timeline, "
            #            "sales process metrics, and supporting documentation.\n\n"
            #            "For additional strategic analysis and insights, see companion document: "
            #            "\"Supplemental Analysis & Insights\"")
            # formatter.add_document_purpose(purpose)

            extraction_result = merger_background_results.get(
                'extraction_result', {})
            strict_summary_result = merger_background_results.get(
                'strict_summary_result', {})

            # Section 1: Chronological Summary
            # doc.add_page_break()
            doc.add_heading('Chronological Summary', 1)

            if strict_summary_result.get('success'):
                formatter.add_simple_numbered_summary(
                    strict_summary_result['summary_text'])

                # Validation info
                if 'validation' in strict_summary_result:
                    formatter.add_horizontal_line()
                    val_para = doc.add_paragraph()
                    val_para.paragraph_format.space_before = Pt(6)

                    val_color = RGBColor(
                        0, 128, 0) if strict_summary_result['validation']['passed'] else RGBColor(255, 0, 0)

                    val_run = val_para.add_run(
                        f"✓ Format Validation: {strict_summary_result['validation']['score']}/{strict_summary_result['validation']['total']} checks passed"
                    )
                    val_run.font.size = Pt(9)
                    val_run.font.color.rgb = val_color

                    if 'max_sentence_length' in strict_summary_result['validation']:
                        val_para.add_run(
                            f" | Max sentence: {strict_summary_result['validation']['max_sentence_length']} words")
            else:
                doc.add_paragraph(
                    f"❌ Error: {strict_summary_result.get('error', 'Unknown error')}")

            # Extract and add other client deliverable sections
            if extraction_result.get('success'):
                doc.add_page_break()
                doc.add_heading(
                    'Extraction of other client deliverable sections', 1)
                extraction_text = extraction_result['extraction_text']
                sections = _parse_extraction_for_document_a(extraction_text)

                for section_title, section_content in sections:
                    doc.add_page_break()
                    doc.add_heading(section_title, 1)

                    # Check if it's a table
                    if _is_table(section_content):
                        formatter.add_professional_table(section_content)
                    else:
                        # Regular text
                        paragraphs = section_content.split('\n\n')
                        for para in paragraphs:
                            if para.strip():
                                clean_para = formatter._clean_markdown(
                                    para.strip())
                                doc.add_paragraph(clean_para)

        # Save the document with custom filename format
        safe_deal_id = deal_id.replace(" ", "_").replace("/", "_")
        docx_filename = f"{safe_deal_id}_summary.docx"
        doc.save(docx_filename)
        logger.info(f"Word document created: {docx_filename}")

        return docx_filename

    except Exception as e:
        error_msg = str(e)
        logger.error(f"Error converting to Word document: {error_msg}")
        raise


def _is_table(text: str) -> bool:
    """Check if text appears to be a table."""
    lines = text.split('\n')
    pipe_lines = [line for line in lines if '|' in line]

    # If more than 50% of lines have pipes, and we have at least 3 lines, it's probably a table
    return len(pipe_lines) >= 3 and len(pipe_lines) / max(len(lines), 1) > 0.5


def _parse_extraction_for_document_a(extraction_text: str) -> List[Tuple[str, str]]:
    """Extract only sections needed for Document A (client deliverables)."""
    required_sections = {
        '2. Starting Point': [],
        '3. Process Structure': [],
        '4. Bidder Universe': [],  # CRITICAL - NOW INCLUDED
        '5. Complete Bid Timeline': [],
        '6. Sales Process Metrics': [],
        '7. Final Round Analysis': [],
        '8. Board Selection Rationale': [],
        '9. Risk Analysis - Regulatory': [],  # CRITICAL - NOW INCLUDED
        '10. Risk Analysis - Financing': [],
        '13. Key Dates Summary': []
    }

    # Map section number → canonical key so LLM variants like
    # "4. Bidder Universe (Complete Census)" or
    # "9. Risk Analysis - Regulatory/Antitrust" still match.
    section_num_map = {}
    for key in required_sections:
        m = re.match(r'^(\d+)\.', key)
        if m:
            section_num_map[m.group(1)] = key

    lines = extraction_text.split('\n')
    current_section = None
    current_content = []

    for line in lines:
        if line.startswith('## '):
            if current_section and current_section in required_sections:
                required_sections[current_section] = '\n'.join(
                    current_content)

            section_name = line.replace('##', '').strip()
            if section_name in required_sections:
                current_section = section_name
                current_content = []
            else:
                # Fall back to number-prefix matching to handle LLM header variations
                m = re.match(r'^(\d+)\.', section_name)
                if m and m.group(1) in section_num_map:
                    current_section = section_num_map[m.group(1)]
                    current_content = []
                else:
                    current_section = None
        elif current_section:
            current_content.append(line)

    if current_section and current_section in required_sections:
        required_sections[current_section] = '\n'.join(current_content)

    return [(k, v) for k, v in required_sections.items() if v]


def _process_content_section(doc: Document, content: str):
    """Helper function to process content and add to document."""
    lines = content.split('\n')
    in_bullet_section = False
    current_header = None

    for line in lines:
        # Check if line is a header (starts with **)
        if line.strip().startswith("**") and line.strip().endswith("**"):
            # Add spacing before new question (except for the first one)
            if current_header:
                doc.add_paragraph()  # Extra spacing between questions
                doc.add_paragraph()  # Double spacing

            # Add header
            header_text = line.strip().replace("**", "")
            header_para = doc.add_paragraph()
            header_run = header_para.add_run(header_text)
            header_run.bold = True
            header_run.font.size = Pt(11)
            # Apply blue color #4f81bd
            header_run.font.color.rgb = RGBColor(0x4f, 0x81, 0xbd)
            header_para.paragraph_format.space_after = Pt(8)
            current_header = header_text
            in_bullet_section = True

        # Check if line is a bullet point (starts with • or -)
        elif line.strip().startswith("•") or line.strip().startswith("-") or line.strip().startswith("*"):
            bullet_text = line.strip()[1:].strip()  # Remove bullet symbol
            if bullet_text:
                para = doc.add_paragraph()

                # Add the + symbol with space
                run = para.add_run("+     ")  # Added extra spaces after +
                run.font.size = Pt(11)

                # Parse and remove ** markers (but don't apply bold)
                parts = re.split(r'(\*\*.*?\*\*)', bullet_text)
                for part in parts:
                    if part.startswith("**") and part.endswith("**"):
                        # Remove ** markers but keep text as regular (not bold)
                        text = part[2:-2]  # Remove ** markers
                        run = para.add_run(text)
                        run.font.size = Pt(11)
                    elif part:
                        # Regular text
                        run = para.add_run(part)
                        run.font.size = Pt(11)

                # Set hanging indent - wrapped lines align with text after "+   "
                para.paragraph_format.left_indent = Pt(25)
                # Hanging indent to pull + back to the left
                para.paragraph_format.first_line_indent = Pt(-20)
                para.paragraph_format.space_after = Pt(6)

        # Regular text
        elif line.strip() and not "="*40 in line:
            # If it looks like it could be a bullet point without a symbol
            if in_bullet_section and len(line.strip()) > 10:
                para = doc.add_paragraph()

                # Add the + symbol with space
                run = para.add_run("+     ")  # Added extra spaces after +
                run.font.size = Pt(11)

                # Parse and remove ** markers (but don't apply bold)
                parts = re.split(r'(\*\*.*?\*\*)', line.strip())
                for part in parts:
                    if part.startswith("**") and part.endswith("**"):
                        # Remove ** markers but keep text as regular (not bold)
                        text = part[2:-2]  # Remove ** markers
                        run = para.add_run(text)
                        run.font.size = Pt(11)
                    elif part:
                        # Regular text
                        run = para.add_run(part)
                        run.font.size = Pt(11)

                # Set hanging indent - wrapped lines align with text after "+   "
                para.paragraph_format.left_indent = Pt(25)
                # Hanging indent to pull + back to the left
                para.paragraph_format.first_line_indent = Pt(-20)
                para.paragraph_format.space_after = Pt(6)
            else:
                # Regular paragraph
                para = doc.add_paragraph(line.strip())
                para.paragraph_format.space_after = Pt(6)


def ask_single_question_on_chunks(processor: QueryProcessor, document_text: str, question: str, question_num: int, max_retries: int = 2) -> Dict[str, Any]:
    """Ask a single question about the document chunks with retry logic."""

    prompt = f"""You are analyzing a merger proxy background section. Please answer the following question based on the document provided.

QUESTION:
{question}

DOCUMENT:
{document_text}

Please provide a specific, detailed answer with dates, names, dollar amounts, and paragraph references where applicable. Be concise but complete."""

    for attempt in range(1, max_retries + 1):
        try:
            print(
                f"  🤔 Asking question {question_num} (attempt {attempt}/{max_retries})...")

            t0 = time.time()
            response = processor.claude_client.messages.create(
                model=processor.claude_model,
                max_tokens=4000,
                messages=[{"role": "user", "content": prompt}],
                temperature=0,
                timeout=180
            )
            elapsed = time.time() - t0

            answer = response.content[0].text.strip()

            print(
                f"  ✅ Answer {question_num} received in {elapsed:.1f}s ({len(answer)} chars)")

            return {
                "question_number": question_num,
                "question": question,
                "answer": answer,
                "response_time_seconds": round(elapsed, 2),
                "attempt": attempt
            }

        except Exception as e:
            error_msg = str(e)
            print(f"  ❌ Error on attempt {attempt}: {error_msg}")
            logger.error(
                f"Error on attempt {attempt} for question {question_num}: {error_msg}")

            if attempt < max_retries:
                wait_time = 5 * attempt
                print(f"  ⏳ Waiting {wait_time} seconds before retry...")
                time.sleep(wait_time)
            else:
                print(
                    f"  ❌ All {max_retries} attempts failed for question {question_num}")
                return {
                    "question_number": question_num,
                    "question": question,
                    "answer": None,
                    "error": error_msg,
                    "attempts": max_retries
                }


def consolidate_qa_answers(processor: QueryProcessor, qa_results: List[Dict[str, Any]]) -> str:
    """Consolidate all individual Q&A results into one cohesive document."""
    print("\n📋 Consolidating all answers into one document...")
    logger.info("Starting consolidation of Q&A answers")

    # Build the Q&A context
    qa_text = ""
    for result in qa_results:
        qa_text += f"\n{'='*80}\n"
        qa_text += f"QUESTION {result['question_number']}: {result['question']}\n"
        qa_text += f"{'='*80}\n"
        qa_text += f"{result.get('answer', 'No answer provided')}\n"

    prompt = f"""You are consolidating individual question-answer pairs about a merger background into one cohesive document.

Here are all the individual Q&A pairs:
{qa_text}

Please create a consolidated document that:
1. Presents all the information from the individual answers
2. Organizes it in a logical, flowing manner
3. Removes any redundancy between answers
4. Maintains all specific details, dates, names, dollar amounts, and references
5. Creates smooth transitions between topics
6. Uses clear section headers to organize the information

The goal is to have one complete, well-organized document that contains everything from the individual answers but reads as a cohesive whole rather than separate Q&As."""

    try:
        t0 = time.time()
        response = processor.claude_client.messages.create(
            model=processor.claude_model,
            max_tokens=8000,
            messages=[{"role": "user", "content": prompt}],
            temperature=0,
            timeout=180
        )
        elapsed = time.time() - t0

        consolidated = response.content[0].text.strip()

        print(
            f"✅ Consolidation completed in {elapsed:.1f}s ({len(consolidated)} chars)\n")
        logger.info(f"Consolidation completed ({len(consolidated)} chars)")
        return consolidated

    except Exception as e:
        error_msg = str(e)
        print(f"❌ Error consolidating answers: {error_msg}")
        logger.error(f"Error consolidating answers: {error_msg}")
        return f"Error consolidating answers: {error_msg}"


def get_background_chunks(processor: QueryProcessor, deal_id: str) -> str:
    """Get background section chunks and format as document text."""
    try:
        print("\n🔍 Searching for background section chunks...")
        logger.info("Searching for background section chunks")

        # Search for background section chunks using title filters
        filters = [
            {"title": {"$eq": "Background of the Mergers"}},
            {"title": {"$eq": "Background of the Transaction"}},
            {"title": {"$eq": "Background of the Merger"}}
        ]

        # Use a dummy vector for filter-only search
        dummy_vector = [0.0] * 3072

        search_response = processor.index.query(
            vector=dummy_vector,
            top_k=8,
            include_metadata=True,
            filter={
                "deal_id": deal_id,
                "$or": filters
            }
        )

        # Format results
        background_chunks = []
        for match in search_response.matches:
            result = {
                'score': match.score,
                'text': match.metadata['original_text'],
                'id': match.id
            }
            background_chunks.append(result)

        print(f"✅ Found {len(background_chunks)} background chunks")
        logger.info(f"Found {len(background_chunks)} background chunks")

        if not background_chunks:
            print("⚠️  No background chunks found")
            return ""

        # Format chunks to document
        document_parts = []
        for i, result in enumerate(background_chunks, 1):
            document_parts.append(f"[Chunk {i}]")
            document_parts.append(result['text'])
            document_parts.append("")

        document_text = "\n".join(document_parts)
        print(f"📄 Document length: {len(document_text):,} characters\n")

        return document_text

    except Exception as e:
        error_msg = str(e)
        logger.error(f"Error getting background chunks: {error_msg}")
        print(f"❌ Error: {error_msg}")
        return ""


def generate_merger_background_analysis(document_text: str) -> Dict[str, Any]:
    """Generate comprehensive background analysis using ProxyBackgroundAnalyzer."""
    try:
        print("\n" + "="*80)
        print("GENERATING MERGER BACKGROUND ANALYSIS")
        print("="*80)
        logger.info("Starting merger background analysis")

        if not document_text:
            print("⚠️  No document text provided")
            return {
                "success": False,
                "error": "No document text provided"
            }

        # Initialize the analyzer
        analyzer = ProxyBackgroundAnalyzer()

        # Stage 1: Deep extraction
        extraction_result = analyzer.stage_1_extraction(document_text)

        if not extraction_result.get('success'):
            return {
                "success": False,
                "error": extraction_result.get('error', 'Extraction failed'),
                "extraction_result": extraction_result
            }

        # Stage 2a: Strict summary
        strict_summary_result = analyzer.stage_2a_strict_summary(
            extraction_result['extraction_text'])

        # Stage 2b: Narrative summary
        # narrative_summary_result = analyzer.stage_2b_narrative_summary(
        #     extraction_result['extraction_text'])

        # Stage 3: Red flags
        # if strict_summary_result.get('success'):
        #     red_flags_result = analyzer.stage_3_red_flags(
        #         extraction_result['extraction_text'],
        #         strict_summary_result['summary_text']
        #     )
        # else:
        #     red_flags_result = {
        #         "red_flags_text": None,
        #         "error": "Strict summary failed",
        #         "success": False
        #     }

        # print(f"✅ Merger background analysis completed")
        # logger.info("Merger background analysis completed")

        return {
            "success": True,
            "extraction_result": extraction_result,
            "strict_summary_result": strict_summary_result,
            # "narrative_summary_result": narrative_summary_result,
            # "red_flags_result": red_flags_result
        }

    except Exception as e:
        error_msg = str(e)
        logger.error(f"Error in merger background analysis: {error_msg}")
        print(f"❌ Error: {error_msg}")
        return {
            "success": False,
            "error": error_msg
        }


def select_deal_id(processor: QueryProcessor) -> str:
    """Helper function to select a deal ID"""
    while True:
        print("\nAvailable deal IDs:")

        # Dictionary mapping display names to actual Pinecone deal_id values
        deal_id_map = {
            "AZEK Company": "AZEK Company",
            "United States Steel Corporation": "United States Steel Corporation 1",
            "Redwood": "Redwood1a1",
            "Regulus": "Regulus14d",
            "Meridianlink": "Meredianlink",
            "Chart": "Chart",
            "Performant": "Performant",
            "City Office": "Cityoffice",
            "Verona": "Verona",
            "Tourmaline": "Tourmaline",
            "Forge": "Forge",
        }

        display_names = list(deal_id_map.keys())

        if not display_names:
            print("No deal IDs found in the index.")
            return input("Please enter a deal ID manually: ")

        for i, display_name in enumerate(display_names, 1):
            print(f"{i}. {display_name}")

        try:
            choice = input(
                "\nEnter the number of the deal ID you want to query (or type the deal ID directly): ")

            # Check if input is a number
            if choice.isdigit() and 1 <= int(choice) <= len(display_names):
                selected_display = display_names[int(choice) - 1]
                return deal_id_map[selected_display]
            # If not a number, check if it's in the map
            elif choice in deal_id_map:
                return deal_id_map[choice]
            # Otherwise, treat as direct deal ID input
            return choice
        except (ValueError, IndexError):
            print("Invalid selection. Please try again.")


def main():
    # Initialize the processor
    processor = QueryProcessor()

    # Select deal ID
    print("\nWelcome to the SEC Filing Query System!")
    deal_id = select_deal_id(processor)
    print(f"\n✅ Selected Deal ID: {deal_id}")

    # Load questions from JSON file
    questions_file = "quetions.json"
    if not os.path.exists(questions_file):
        print(f"❌ Questions file not found: {questions_file}")
        logger.error(f"Questions file not found: {questions_file}")
        return

    try:
        with open(questions_file, "r", encoding="utf-8") as f:
            questions_data = json.load(f)
        logger.info(
            f"Loaded {len(questions_data)} questions from {questions_file}")
        print(
            f"📋 Loaded {len(questions_data)} questions from {questions_file}")
    except Exception as e:
        print(f"❌ Error loading questions file: {str(e)}")
        logger.error(f"Error loading questions file: {str(e)}")
        return

    # Prepare to collect all results
    all_results = []
    safe_deal_id = deal_id.replace(" ", "_").replace("/", "_")
    merged_filename = f"{safe_deal_id}_arb_summary.txt"

    print("\n" + "="*80)
    print("PROCESSING ALL QUESTIONS")
    print("="*80 + "\n")

    # Process each question
    for question_key, question in questions_data.items():
        print(f"\n{'='*80}")
        print(f"Processing: {question_key}")
        print(f"Question: {question}")
        print("="*80)
        logger.info(f"Processing {question_key}: {question}")

        try:
            # Search for similar chunks
            print("🔍 Searching for relevant chunks...")
            results = processor.search_similar_chunks(question, deal_id)
            print(f"✅ Found {len(results)} relevant chunks")

            # Get Claude response
            print("🤔 Getting Claude response...")
            claude_answer = processor.get_claude_response(question, results)
            print(f"✅ Received answer ({len(claude_answer)} chars)")

            # Generate arbitrage summary
            arb_summary = processor.generate_arb_summary(
                claude_answer, question)

            # Store result
            result_entry = {
                'question_key': question_key,
                'question': question,
                'answer': claude_answer,
                'summary': arb_summary
            }
            all_results.append(result_entry)

            # Display summary preview
            print("\n📊 SUMMARY PREVIEW:")
            print("-" * 80)
            summary_preview = arb_summary[:400] + \
                ("..." if len(arb_summary) > 400 else "")
            print(summary_preview)
            print("-" * 80)

        except Exception as e:
            error_msg = str(e)
            logger.error(f"Error processing {question_key}: {error_msg}")
            print(f"❌ Error: {error_msg}")

            # Store error result
            result_entry = {
                'question_key': question_key,
                'question': question,
                'answer': f"Error: {error_msg}",
                'summary': f"Error generating summary: {error_msg}"
            }
            all_results.append(result_entry)

    # Save merged results to file
    print("\n" + "="*80)
    print("SAVING MERGED RESULTS")
    print("="*80)

    try:
        with open(merged_filename, "w", encoding="utf-8") as f:
            # Write each question's summary only (no header)
            for i, result in enumerate(all_results, 1):
                f.write(result['summary'] + "\n\n")

        logger.info(f"Merged summary saved to {merged_filename}")
        print(f"\n✅ Merged summary saved to: {merged_filename}")
        print(f"   Total questions processed: {len(all_results)}")
        print(f"   File size: {os.path.getsize(merged_filename):,} bytes")

    except Exception as e:
        error_msg = str(e)
        logger.error(f"Error saving merged file: {error_msg}")
        print(f"❌ Error saving merged file: {error_msg}")

    # Generate merger background analysis
    print("\n" + "="*80)
    print("GENERATING MERGER BACKGROUND ANALYSIS")
    print("="*80)

    merger_background_results = None
    # breakpoint()
    try:
        # Get background chunks
        document_text = get_background_chunks(processor, deal_id)

        if document_text:
            # Generate merger background analysis
            merger_background_results = generate_merger_background_analysis(
                document_text)

            if merger_background_results.get('success'):
                print(f"✅ Merger background analysis completed successfully")
                logger.info(
                    "Merger background analysis completed successfully")
            else:
                print(
                    f"⚠️  Merger background analysis failed: {merger_background_results.get('error', 'Unknown error')}")
                logger.warning(
                    f"Merger background analysis failed: {merger_background_results.get('error', 'Unknown error')}")
        else:
            print(f"⚠️  No background chunks found, skipping merger background analysis")
            logger.warning(
                "No background chunks found, skipping merger background analysis")

    except Exception as e:
        error_msg = str(e)
        logger.error(
            f"Error generating merger background analysis: {error_msg}")
        print(f"❌ Error generating merger background analysis: {error_msg}")

    # Convert to Word document (combining Q&A and merger background analysis)
    print("\n" + "="*80)
    print("CONVERTING TO WORD DOCUMENT")
    print("="*80)

    try:
        docx_filename = convert_txt_to_docx(
            merged_filename, deal_id, merger_background_results)
        print(f"\n✅ Word document created: {docx_filename}")
        print(f"   File size: {os.path.getsize(docx_filename):,} bytes")
    except Exception as e:
        error_msg = str(e)
        logger.error(f"Error creating Word document: {error_msg}")
        print(f"❌ Error creating Word document: {error_msg}")

    print("\n" + "="*80)
    print("✅ PROCESSING COMPLETE!")
    print("="*80)


if __name__ == "__main__":
    main()
