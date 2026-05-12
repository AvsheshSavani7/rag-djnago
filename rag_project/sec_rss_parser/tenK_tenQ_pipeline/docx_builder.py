"""DOCX generation: single-filing reports and multi-filing comparison reports."""

import json
import re
from datetime import datetime
from pathlib import Path
from typing import List

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from .config import (
    COLOR_TIMING, COLOR_REGULATORY, COLOR_BOTH,
    COLOR_SIGNIFICANT, COLOR_MODERATE, COLOR_MINOR,
    COLOR_NEW, COLOR_UNCHANGED,
)
from .models import DealContext


# =============================================================================
# SHARED DOCX HELPERS
# =============================================================================

def _apply_shading(run, hex_color: str) -> None:
    rPr = run._r.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    rPr.append(shd)


def _shade_cell(cell, hex_color: str) -> None:
    """Apply background fill color to a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def _make_doc_base(deal: DealContext, filing_labels_str: str) -> Document:
    """Create a Document with standard header."""
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(10)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15

    hdr = doc.add_paragraph()
    r = hdr.add_run(f"{deal.acquirer_company} / {deal.target_company} ({deal.ticker})")
    r.font.size = Pt(13); r.bold = True

    fl = doc.add_paragraph()
    r = fl.add_run(filing_labels_str)
    r.font.size = Pt(9); r.font.color.rgb = RGBColor(80, 80, 80)

    doc.add_paragraph()
    return doc


# =============================================================================
# CHANGE REPORT CONSTANTS AND HELPERS
# =============================================================================

_SEVERITY_GROUPS = {
    "significant": ("SIGNIFICANT CHANGES", COLOR_SIGNIFICANT, RGBColor(183, 28, 28)),
    "moderate":    ("MODERATE CHANGES",    COLOR_MODERATE,    RGBColor(230, 81, 0)),
    "minor":       ("MINOR CHANGES",       COLOR_MINOR,       RGBColor(156, 110, 0)),
    "new":         ("NEW DISCLOSURES",     COLOR_NEW,         RGBColor(27, 94, 32)),
}

_PASS_COLORS = {
    "timing":         (RGBColor(21, 101, 192), "TIMING"),
    "regulatory":     (RGBColor(173, 20, 87),  "REGULATORY"),
    "legal_language": (RGBColor(156, 110, 0),  "LEGAL LANGUAGE"),
}

_TIER1_SECTION_PATTERNS = [
    re.compile(r'\bnote\s+\d', re.IGNORECASE),
    re.compile(r'financial statements', re.IGNORECASE),
    re.compile(r'risk\s+factors?', re.IGNORECASE),
    re.compile(r'\bitem\s+1a\b', re.IGNORECASE),
]
_TIER2_SECTION_PATTERNS = [
    re.compile(r"management.s discussion", re.IGNORECASE),
    re.compile(r'\bmd&?a\b', re.IGNORECASE),
    re.compile(r'\bitem\s+2\b', re.IGNORECASE),
]
_RISK_FACTOR_PAT = re.compile(r'risk\s+factors?|item\s+1a', re.IGNORECASE)


def _classify_section_tier(section: str, text: str = "") -> int:
    """1 = merger note / risk factors, 2 = MD&A, 3 = other."""
    for pat in _TIER1_SECTION_PATTERNS:
        if pat.search(section):
            return 1
    for pat in _TIER2_SECTION_PATTERNS:
        if pat.search(section):
            return 2
    return 3


def _section_tier_label(section: str) -> str:
    if _RISK_FACTOR_PAT.search(section):
        return "RISK FACTORS"
    for pat in _TIER1_SECTION_PATTERNS:
        if pat.search(section):
            return "MERGER NOTE"
    for pat in _TIER2_SECTION_PATTERNS:
        if pat.search(section):
            return "MD&A"
    return "OTHER"


def _pass_priority(result: dict) -> int:
    """0 = timing/regulatory fired (highest), 1 = legal_language only."""
    for pk in ["timing", "regulatory"]:
        if (result.get(pk) or {}).get("changed"):
            return 0
    return 1


def _render_group_header(doc, label: str, bg_hex: str, fg_color) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(f"  {label}  ")
    _apply_shading(run, bg_hex)
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = fg_color


def _render_item_header(doc, num: int, result: dict, sev_key: str) -> None:
    sev_label = "NEW" if sev_key == "new" else result.get("overall_severity", "none").upper()
    bg_hex = {
        "significant": COLOR_SIGNIFICANT, "moderate": COLOR_MODERATE,
        "minor": COLOR_MINOR, "new": COLOR_NEW,
    }.get(sev_key, COLOR_UNCHANGED)
    txt_color = _SEVERITY_GROUPS.get(sev_key, ("", "", RGBColor(80, 80, 80)))[2]

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    tag = p.add_run(f" {num}. [{sev_label}] ")
    _apply_shading(tag, bg_hex)
    tag.bold = True
    tag.font.size = Pt(9)
    tag.font.color.rgb = txt_color

    section = result.get("section", "")[:60]
    if section:
        sec_run = p.add_run(f"  {section}")
        sec_run.font.size = Pt(8)
        sec_run.font.color.rgb = RGBColor(120, 120, 120)


def _resolve_prior_label(result: dict, prior_labels: list) -> str:
    for pk in ["timing", "regulatory", "legal_language"]:
        finding = result.get(pk) or {}
        if "_prior_label" in finding:
            return finding["_prior_label"]
    return ", ".join(prior_labels)


# =============================================================================
# SINGLE-FILING REPORTS
# =============================================================================

def generate_single_filing_report_fulsome(
    excerpts: list,
    deal: DealContext,
    filing_label: str,
    output_path: Path,
) -> None:
    """Generate a fulsome DOCX for a single filing with all commentary."""
    priority_categories = {"risk", "deal_terms", "timeline", "regulatory"}

    filtered = [
        e for e in excerpts
        if (e.get("timing_flag") or e.get("regulatory_flag"))
        and e.get("category", "general") in priority_categories
    ]

    if not filtered:
        filtered = [e for e in excerpts if e.get("timing_flag") or e.get("regulatory_flag")]

    def sort_key(e):
        cat = e.get("category", "general")
        if cat == "risk": return (0, -e.get("relevance_score", 0))
        elif cat == "deal_terms": return (1, -e.get("relevance_score", 0))
        else: return (2, -e.get("relevance_score", 0))

    filtered.sort(key=sort_key)

    doc = _make_doc_base(deal, f"Filing: {filing_label}")
    doc.add_heading('Key Excerpts', 1)

    for idx, excerpt in enumerate(filtered, 1):
        timing = excerpt.get("timing_flag", False)
        regulatory = excerpt.get("regulatory_flag", False)
        category = excerpt.get("category", "general")
        score = excerpt.get("relevance_score", 0)

        if timing and regulatory:
            bg_color = COLOR_BOTH
            flag_label = "TIMING + REGULATORY"
        elif timing:
            bg_color = COLOR_TIMING
            flag_label = "TIMING"
        else:
            bg_color = COLOR_REGULATORY
            flag_label = "REGULATORY"

        header_p = doc.add_paragraph()
        tag = header_p.add_run(f" {idx}. [{flag_label}] ")
        _apply_shading(tag, bg_color)
        tag.bold = True; tag.font.size = Pt(9)
        tag.font.color.rgb = (
            RGBColor(0, 0, 0) if timing and regulatory
            else RGBColor(21, 101, 192) if timing
            else RGBColor(173, 20, 87)
        )

        cat_score = header_p.add_run(f"  {category.upper()} (Score: {score}/10)")
        cat_score.font.size = Pt(8); cat_score.font.color.rgb = RGBColor(120, 120, 120)

        sec_p = doc.add_paragraph()
        sec_p.paragraph_format.left_indent = Inches(0.15)
        sec_r = sec_p.add_run(f"Section: {excerpt.get('section', 'Unknown')}")
        sec_r.font.size = Pt(8); sec_r.italic = True; sec_r.font.color.rgb = RGBColor(100, 100, 100)

        if timing and excerpt.get("timing_assessment"):
            assess_p = doc.add_paragraph()
            assess_p.paragraph_format.left_indent = Inches(0.15)
            assess_r = assess_p.add_run(f"Timing: {excerpt['timing_assessment']}")
            assess_r.font.size = Pt(9); assess_r.font.color.rgb = RGBColor(21, 101, 192)

        if regulatory and excerpt.get("regulatory_assessment"):
            assess_p = doc.add_paragraph()
            assess_p.paragraph_format.left_indent = Inches(0.15)
            assess_r = assess_p.add_run(f"Regulatory: {excerpt['regulatory_assessment']}")
            assess_r.font.size = Pt(9); assess_r.font.color.rgb = RGBColor(173, 20, 87)

        text_p = doc.add_paragraph()
        text_p.paragraph_format.left_indent = Inches(0.25)
        text_p.paragraph_format.space_before = Pt(4)
        clean = excerpt.get("text", "").replace('\n', ' ').replace('  ', ' ').strip()
        text_r = text_p.add_run(clean)
        text_r.font.size = Pt(9)

        doc.add_paragraph()

    doc.save(output_path)
    print(f"  Fulsome report saved: {output_path}")


def generate_single_filing_report_concise(
    excerpts: list,
    deal: DealContext,
    filing_label: str,
    output_path: Path,
) -> None:
    """Generate a concise DOCX for a single filing — no commentary or scores."""
    priority_categories = {"risk", "deal_terms", "timeline", "regulatory"}

    filtered = [
        e for e in excerpts
        if (e.get("timing_flag") or e.get("regulatory_flag"))
        and e.get("category", "general") in priority_categories
    ]

    if not filtered:
        filtered = [e for e in excerpts if e.get("timing_flag") or e.get("regulatory_flag")]

    def sort_key(e):
        cat = e.get("category", "general")
        if cat == "risk": return (0, -e.get("relevance_score", 0))
        elif cat == "deal_terms": return (1, -e.get("relevance_score", 0))
        else: return (2, -e.get("relevance_score", 0))

    filtered.sort(key=sort_key)

    doc = _make_doc_base(deal, f"Filing: {filing_label}")
    doc.add_heading('Key Excerpts', 1)

    for idx, excerpt in enumerate(filtered, 1):
        timing = excerpt.get("timing_flag", False)
        regulatory = excerpt.get("regulatory_flag", False)
        category = excerpt.get("category", "general")

        if timing and regulatory:
            bg_color = COLOR_BOTH
            flag_label = "TIMING + REGULATORY"
        elif timing:
            bg_color = COLOR_TIMING
            flag_label = "TIMING"
        else:
            bg_color = COLOR_REGULATORY
            flag_label = "REGULATORY"

        header_p = doc.add_paragraph()
        tag = header_p.add_run(f" {idx}. [{flag_label}] ")
        _apply_shading(tag, bg_color)
        tag.bold = True; tag.font.size = Pt(9)
        tag.font.color.rgb = (
            RGBColor(0, 0, 0) if timing and regulatory
            else RGBColor(21, 101, 192) if timing
            else RGBColor(173, 20, 87)
        )

        cat_label = header_p.add_run(f"  {category.upper()}")
        cat_label.font.size = Pt(8); cat_label.font.color.rgb = RGBColor(120, 120, 120)

        sec_p = doc.add_paragraph()
        sec_p.paragraph_format.left_indent = Inches(0.15)
        sec_r = sec_p.add_run(f"Section: {excerpt.get('section', 'Unknown')}")
        sec_r.font.size = Pt(8); sec_r.italic = True; sec_r.font.color.rgb = RGBColor(100, 100, 100)

        text_p = doc.add_paragraph()
        text_p.paragraph_format.left_indent = Inches(0.25)
        text_p.paragraph_format.space_before = Pt(4)
        clean = excerpt.get("text", "").replace('\n', ' ').replace('  ', ' ').strip()
        text_r = text_p.add_run(clean)
        text_r.font.size = Pt(9)

        doc.add_paragraph()

    doc.save(output_path)
    print(f"  Concise report saved: {output_path}")


# =============================================================================
# MULTI-FILING COMPARISON REPORTS
# =============================================================================

def generate_change_report(
    all_comparison_steps: list,
    deal: DealContext,
    filing_labels: List[str],
    output_path: Path,
) -> None:
    """
    Generate a change report DOCX — what changed between filings.
    Findings grouped by severity, each with current/prior text quotes and factual analysis.
    """
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(10)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15

    hdr = doc.add_paragraph()
    r = hdr.add_run(f"{deal.acquirer_company} / {deal.target_company} ({deal.ticker})")
    r.font.size = Pt(13)
    r.bold = True

    fl = doc.add_paragraph()
    r = fl.add_run(' → '.join(filing_labels))
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(80, 80, 80)

    doc.add_paragraph()

    for step_idx, step in enumerate(all_comparison_steps):
        current_label = step["current_label"]
        prior_labels = step["prior_labels"]
        merged = step["merged_results"]

        if len(prior_labels) == 1:
            heading_text = f"{current_label}  vs  {prior_labels[0]}"
        else:
            primary = prior_labels[-1]
            context = ", ".join(prior_labels[:-1])
            heading_text = f"{current_label}  vs  {primary}  (with {context})"
        doc.add_heading(heading_text, 1)

        if not merged:
            p = doc.add_paragraph()
            r = p.add_run("No material changes detected.")
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(100, 100, 100)
            doc.add_paragraph()
            continue

        sig = sum(1 for r in merged if r["overall_severity"] == "significant")
        mod = sum(1 for r in merged if r["overall_severity"] == "moderate")
        minor_ct = sum(1 for r in merged if r["overall_severity"] == "minor")
        new_ct = sum(1 for r in merged if r.get("is_new", False))

        stat = doc.add_paragraph()
        stat.paragraph_format.space_after = Pt(8)
        parts = []
        if sig: parts.append(f"{sig} significant")
        if mod: parts.append(f"{mod} moderate")
        if minor_ct: parts.append(f"{minor_ct} minor")
        if new_ct: parts.append(f"{new_ct} new disclosure(s)")
        sr = stat.add_run(f"{len(merged)} change(s) found: {', '.join(parts)}")
        sr.font.size = Pt(9)
        sr.bold = True

        # Boilerplate detection for "mostly new" comparisons
        total_ct = len(merged)
        new_disc_ct = sum(1 for r in merged if r.get("is_new", False))
        mostly_new = total_ct > 0 and new_disc_ct > total_ct * 0.7

        if mostly_new:
            _substance_patterns = [
                re.compile(r'\$[\d,.]+\s*(million|billion|per share|in cash)', re.IGNORECASE),
                re.compile(r'\$\s*\d+\s', re.IGNORECASE),
                re.compile(r'(january|february|march|april|may|june|july|august|september|october|november|december)\s+\d{1,2},?\s+\d{4}', re.IGNORECASE),
                re.compile(r'\b(HSR|antitrust|money transmitter|CFIUS|FTC|DOJ)\b', re.IGNORECASE),
                re.compile(r'\b(termination fee|break.?up fee|reverse termination)\b', re.IGNORECASE),
                re.compile(r'\b(per share|consideration of|exchange ratio)\b', re.IGNORECASE),
                re.compile(r'\b(waiver|consent solicitation|supplemental indenture)\b', re.IGNORECASE),
                re.compile(r'\b(downgrad|rating|Fitch|Moody|S&P|AM Best)\b', re.IGNORECASE),
                re.compile(r'\b(complaint|lawsuit|injunction|demand letter)\b', re.IGNORECASE),
                re.compile(r'\b(transaction costs?|acquisition.related costs?|deal (costs?|expenses?))\b', re.IGNORECASE),
                re.compile(r'\b(EBITDA|adjusted ebitda|non.?GAAP)\b', re.IGNORECASE),
                re.compile(r'\b(indebtedness|debt incurrence|debt covenant)\b', re.IGNORECASE),
            ]
            _boilerplate_markers = [
                re.compile(r'forward.looking\s+statement', re.IGNORECASE),
                re.compile(r'safe\s+harbor', re.IGNORECASE),
                re.compile(r'^exhibit\s+(index|listing)', re.IGNORECASE),
            ]

            def _has_substance(result_item):
                text = (result_item.get("text") or "")
                analysis_text = ""
                for pk in ["timing", "regulatory", "legal_language"]:
                    a = ((result_item.get(pk) or {}).get("analysis") or "")
                    analysis_text += " " + a
                combined = text + " " + analysis_text
                if any(bp.search(combined) for bp in _boilerplate_markers):
                    if not any(sp.search(combined) for sp in _substance_patterns):
                        return False
                substance_hits = sum(1 for sp in _substance_patterns if sp.search(combined))
                return substance_hits >= 1

            for result in merged:
                if result.get("is_new", False) and result["overall_severity"] in ("significant", "moderate"):
                    if not _has_substance(result):
                        result["_original_severity"] = result["overall_severity"]
                        result["overall_severity"] = "minor"

        severity_groups = {
            "significant": [], "moderate": [], "minor": [], "new": [],
        }
        for result in merged:
            if result.get("is_new") and result["overall_severity"] == "none":
                severity_groups["new"].append(result)
            elif result["overall_severity"] in severity_groups:
                severity_groups[result["overall_severity"]].append(result)
            else:
                if any((result.get(pk) or {}).get("changed") for pk in ["timing", "regulatory", "legal_language"]):
                    severity_groups["minor"].append(result)

        for sev_key in severity_groups:
            severity_groups[sev_key].sort(
                key=lambda r: (
                    _classify_section_tier(r.get("section", ""), r.get("text", "")),
                    _pass_priority(r),
                )
            )

        item_num = 0
        for sev_key in ["significant", "moderate", "minor", "new"]:
            items = severity_groups[sev_key]
            if not items:
                continue

            group_label, group_bg, group_fg = _SEVERITY_GROUPS[sev_key]
            _render_group_header(doc, group_label, group_bg, group_fg)

            for result in items:
                item_num += 1
                is_new = result.get("is_new", False)

                active_passes = [
                    pk for pk in ["timing", "regulatory", "legal_language"]
                    if (result.get(pk) or {}).get("changed")
                ]

                _render_item_header(doc, item_num, result, sev_key)

                _skip_starts = {"this is", "this represents", "this reflects",
                                "significantly", "importantly", "notably",
                                "critically", "most notably", "it reveals",
                                "this was not"}
                _skip_contains = {"is notable", "is a deliberate", "is significant",
                                  "is a material", "a lawyer may", "a lawyer would",
                                  "may have done", "may have made",
                                  "minor punctuation", "represents a significant",
                                  "significantly expanded"}
                _novelty_prefixes = [
                    "new disclosure ", "new disclosure of ", "new disclosure establishes ",
                    "new disclosure reaffirms ", "new disclosure confirming ",
                    "new risk factor ", "new risk factor disclosure ",
                    "new risk factor header ", "new risk factor paragraph ",
                    "new paragraph ", "new exhibit ",
                    "new md&a ", "new md&a paragraph ",
                    "new note ", "new forward-looking statement ",
                    "entirely new ", "entirely new paragraph ",
                    "entirely new note ", "entirely new section ",
                    "new non-gaap ",
                    "first-time disclosure of ",
                    "continuation of ",
                    "standalone exhibit reference to ",
                ]

                def _is_factual(sent_text):
                    low = sent_text.lower()
                    if any(low.startswith(s) for s in _skip_starts):
                        return False
                    if any(phrase in low for phrase in _skip_contains):
                        return False
                    return True

                def _strip_novelty(sent_text):
                    result_text = sent_text
                    low = result_text.lower()
                    for prefix in _novelty_prefixes:
                        if low.startswith(prefix):
                            result_text = result_text[len(prefix):]
                            low = result_text.lower()
                            break
                    _leftover = ["of ", "that ", "the ", "header ", "listing ",
                                 "reference ", "paragraph ", "section ",
                                 "confirming ", "disclosing ", "establishing "]
                    for lo in _leftover:
                        if low.startswith(lo):
                            result_text = result_text[len(lo):]
                            break
                    if len(result_text) > 20 and result_text != sent_text:
                        return result_text[0].upper() + result_text[1:]
                    return sent_text

                def _clean_sentence(sent_text):
                    cleaned = _strip_novelty(sent_text)
                    if _is_factual(cleaned):
                        return cleaned
                    return ""

                def _first_factual(text):
                    for sent in text.split(". "):
                        sent = sent.strip().rstrip(".")
                        if not sent or len(sent) < 15:
                            continue
                        result_sent = _clean_sentence(sent)
                        if result_sent:
                            return result_sent.rstrip(".") + "."
                    return ""

                summary_sentence = ""
                for pk in ["timing", "regulatory"]:
                    if pk not in active_passes:
                        continue
                    a = ((result.get(pk) or {}).get("analysis") or "").strip()
                    if a:
                        summary_sentence = _first_factual(a)
                    if summary_sentence:
                        break

                if not summary_sentence:
                    ll_f = result.get("legal_language") or {}
                    nc_list = ll_f.get("notable_changes") or []

                    def _headline_score(nc):
                        old_p = nc.get("old_phrase", "") or ""
                        new_p = nc.get("new_phrase", "") or ""
                        interp_l = (nc.get("interpretation", "") or "").lower()
                        s = len(old_p) + len(new_p)
                        if any(t in interp_l for t in ["note number", "numbering", "formatting",
                                "defined term", "stylistic", "reorder", "renumber"]):
                            s *= 0.1
                        if any(k in interp_l for k in ["filing", "approval", "settlement", "closing",
                                "timeline", "regulatory", "hearing", "clearance", "expired",
                                "antitrust", "ferc", "mpsc", "hsr", "sec ", "s-4", "proxy"]):
                            s *= 3.0
                        if old_p and new_p:
                            _strip_h = lambda t: set(re.sub(r'[^\w\s]', '', t.lower()).split())
                            ow = _strip_h(old_p)
                            nw = _strip_h(new_p)
                            if ow and nw and len(ow & nw) / max(len(ow), len(nw)) > 0.85:
                                s *= 0.05
                        if old_p and not new_p:
                            s *= 0.2
                        return s

                    nc_sorted = sorted(nc_list, key=_headline_score, reverse=True)
                    for nc in nc_sorted[:3]:
                        interp = (nc.get("interpretation") or "").strip()
                        if interp and len(interp) > 15:
                            candidate = _first_factual(interp)
                            if candidate:
                                summary_sentence = candidate
                                break

                if not summary_sentence and "legal_language" in active_passes:
                    a = ((result.get("legal_language") or {}).get("analysis") or "").strip()
                    if a:
                        summary_sentence = _first_factual(a)

                _MAX_HEADLINE = 350
                if summary_sentence and len(summary_sentence) > _MAX_HEADLINE:
                    truncated = summary_sentence[:_MAX_HEADLINE].rsplit(" ", 1)[0]
                    summary_sentence = truncated.rstrip(".,;:—") + "..."

                if summary_sentence:
                    sp = doc.add_paragraph()
                    sp.paragraph_format.left_indent = Inches(0.15)
                    sp.paragraph_format.space_before = Pt(2)
                    sp.paragraph_format.space_after = Pt(4)
                    sr = sp.add_run(summary_sentence)
                    sr.font.size = Pt(10)
                    sr.bold = True

                prior_lbl = _resolve_prior_label(result, prior_labels)
                ll = result.get("legal_language") or {}
                notable = ll.get("notable_changes") or []
                has_phrase_quotes = any(
                    (c.get("old_phrase") or "").strip() or (c.get("new_phrase") or "").strip()
                    for c in notable
                )

                def _score_notable_change(nc: dict) -> float:
                    old_p = nc.get("old_phrase", "") or ""
                    new_p = nc.get("new_phrase", "") or ""
                    interp = (nc.get("interpretation", "") or "").lower()
                    score = len(old_p) + len(new_p)
                    for t in ["note number", "numbering", "formatting", "defined term",
                              "stylistic", "reorder", "renumber"]:
                        if t in interp:
                            score *= 0.1
                            break
                    for s in ["filing", "approval", "settlement", "closing", "timeline",
                              "regulatory", "hearing", "clearance", "expired", "antitrust",
                              "ferc", "mpsc", "hsr", "sec ", "s-4", "proxy"]:
                        if s in interp:
                            score *= 3.0
                            break
                    if old_p and new_p:
                        _strip_fn = lambda s: set(re.sub(r'[^\w\s]', '', s.lower()).split())
                        old_words = _strip_fn(old_p)
                        new_words = _strip_fn(new_p)
                        if old_words and new_words:
                            overlap = len(old_words & new_words) / max(len(old_words), len(new_words))
                            if overlap > 0.85:
                                score *= 0.05
                    if old_p and not new_p:
                        score *= 0.2
                    return score

                if has_phrase_quotes:
                    ranked = sorted(notable, key=_score_notable_change, reverse=True)
                    for change in ranked[:2]:
                        old_p = (change.get("old_phrase") or "").strip()
                        new_p = (change.get("new_phrase") or "").strip()
                        if not old_p and not new_p:
                            continue
                        if new_p:
                            cp = doc.add_paragraph()
                            cp.paragraph_format.left_indent = Inches(0.25)
                            cp.paragraph_format.space_before = Pt(2)
                            cp.paragraph_format.space_after = Pt(1)
                            cl = cp.add_run(f"{current_label}: ")
                            cl.bold = True; cl.font.size = Pt(8)
                            cl.font.color.rgb = RGBColor(140, 140, 140)
                            ct = cp.add_run(f'"{new_p}"')
                            ct.font.size = Pt(8)
                            ct.font.color.rgb = RGBColor(80, 80, 80)
                        if old_p:
                            pp = doc.add_paragraph()
                            pp.paragraph_format.left_indent = Inches(0.25)
                            pp.paragraph_format.space_before = Pt(0)
                            pp.paragraph_format.space_after = Pt(3)
                            pl = pp.add_run(f"{prior_lbl}: ")
                            pl.bold = True; pl.font.size = Pt(8)
                            pl.font.color.rgb = RGBColor(140, 140, 140)
                            pt = pp.add_run(f'"{old_p}"')
                            pt.font.size = Pt(8)
                            pt.font.color.rgb = RGBColor(80, 80, 80)
                        else:
                            pp = doc.add_paragraph()
                            pp.paragraph_format.left_indent = Inches(0.25)
                            pp.paragraph_format.space_before = Pt(0)
                            pp.paragraph_format.space_after = Pt(3)
                            pl = pp.add_run(f"{prior_lbl}: ")
                            pl.bold = True; pl.font.size = Pt(8)
                            pl.font.color.rgb = RGBColor(140, 140, 140)
                            pt = pp.add_run("Not present")
                            pt.font.size = Pt(8)
                            pt.font.italic = True
                            pt.font.color.rgb = RGBColor(160, 160, 160)
                else:
                    current_text = (result.get("text") or "").replace('\n', ' ').replace('  ', ' ').strip()
                    prior_text = (result.get("prior_text") or "").replace('\n', ' ').replace('  ', ' ').strip()
                    _QUOTE_MAX = 300
                    if current_text:
                        ct_trunc = current_text[:_QUOTE_MAX] + ("..." if len(current_text) > _QUOTE_MAX else "")
                        cp = doc.add_paragraph()
                        cp.paragraph_format.left_indent = Inches(0.25)
                        cp.paragraph_format.space_before = Pt(2)
                        cp.paragraph_format.space_after = Pt(1)
                        cl = cp.add_run(f"{current_label}: ")
                        cl.bold = True; cl.font.size = Pt(8)
                        cl.font.color.rgb = RGBColor(140, 140, 140)
                        ct = cp.add_run(f'"{ct_trunc}"')
                        ct.font.size = Pt(8)
                        ct.font.color.rgb = RGBColor(80, 80, 80)
                    if prior_text:
                        pt_trunc = prior_text[:_QUOTE_MAX] + ("..." if len(prior_text) > _QUOTE_MAX else "")
                        pp = doc.add_paragraph()
                        pp.paragraph_format.left_indent = Inches(0.25)
                        pp.paragraph_format.space_before = Pt(0)
                        pp.paragraph_format.space_after = Pt(3)
                        pl = pp.add_run(f"{prior_lbl}: ")
                        pl.bold = True; pl.font.size = Pt(8)
                        pl.font.color.rgb = RGBColor(140, 140, 140)
                        pt = pp.add_run(f'"{pt_trunc}"')
                        pt.font.size = Pt(8)
                        pt.font.color.rgb = RGBColor(80, 80, 80)
                    elif current_text and not prior_text:
                        pp = doc.add_paragraph()
                        pp.paragraph_format.left_indent = Inches(0.25)
                        pp.paragraph_format.space_before = Pt(0)
                        pp.paragraph_format.space_after = Pt(3)
                        pl = pp.add_run(f"{prior_lbl}: ")
                        pl.bold = True; pl.font.size = Pt(8)
                        pl.font.color.rgb = RGBColor(140, 140, 140)
                        pt = pp.add_run("Not present")
                        pt.font.size = Pt(8)
                        pt.font.italic = True
                        pt.font.color.rgb = RGBColor(160, 160, 160)

                meta_parts = []
                section = result.get("section", "")
                tier_label = _section_tier_label(section)
                if tier_label != "OTHER":
                    meta_parts.append(f"[{tier_label}]")
                if active_passes:
                    tag_labels = [_PASS_COLORS[pk][1].title() for pk in active_passes]
                    meta_parts.append(", ".join(tag_labels))
                if section:
                    meta_parts.append(section)
                if meta_parts:
                    mp = doc.add_paragraph()
                    mp.paragraph_format.left_indent = Inches(0.15)
                    mp.paragraph_format.space_before = Pt(2)
                    mp.paragraph_format.space_after = Pt(0)
                    mr = mp.add_run(" | ".join(meta_parts))
                    mr.font.size = Pt(7)
                    mr.font.italic = True
                    mr.font.color.rgb = RGBColor(170, 170, 170)

                sep = doc.add_paragraph()
                sep.paragraph_format.space_before = Pt(2)
                sep.paragraph_format.space_after = Pt(2)

        if step_idx < len(all_comparison_steps) - 1:
            doc.add_page_break()

    doc.save(output_path)
    print(f"  Change report saved: {output_path}")


def generate_raw_text_report(
    all_comparison_steps: list,
    deal: DealContext,
    filing_labels: List[str],
    output_path: Path,
) -> None:
    """Generate a DOCX containing only the raw filing text excerpts."""
    doc = _make_doc_base(deal, ' → '.join(filing_labels))

    for step_idx, step in enumerate(all_comparison_steps):
        current_label = step["current_label"]
        prior_labels = step["prior_labels"]
        merged = step["merged_results"]

        doc.add_heading(f"{current_label}  vs  {', '.join(prior_labels)}", 1)

        if not merged:
            np = doc.add_paragraph()
            np.add_run("No excerpts.").font.size = Pt(10)
            np.runs[0].font.color.rgb = RGBColor(100, 100, 100)
            doc.add_paragraph()
            continue

        for idx, result in enumerate(merged, 1):
            section = result.get("section", "")
            clause_text = result.get("text", "").replace('\n', ' ').replace('  ', ' ').strip()

            sec_p = doc.add_paragraph()
            sr = sec_p.add_run(f"{idx}. {section}")
            sr.bold = True
            sr.font.size = Pt(9)
            sr.font.color.rgb = RGBColor(60, 60, 60)

            if clause_text:
                text_p = doc.add_paragraph()
                text_p.paragraph_format.left_indent = Inches(0.25)
                tr = text_p.add_run(clause_text)
                tr.font.size = Pt(9)
                tr.font.color.rgb = RGBColor(30, 30, 30)

            sep = doc.add_paragraph()
            sep.paragraph_format.space_before = Pt(2)
            sep.paragraph_format.space_after = Pt(2)

        if step_idx < len(all_comparison_steps) - 1:
            doc.add_page_break()

    doc.save(output_path)
    print(f"  Raw text report saved: {output_path}")


def generate_full_comparison_json(
    all_comparison_steps: list,
    deal: DealContext,
    filing_labels: List[str],
    output_path: Path,
) -> None:
    output = {
        "metadata": {
            "ticker": deal.ticker,
            "target": deal.target_company,
            "acquirer": deal.acquirer_company,
            "deal_value": deal.deal_value,
            "filings_analyzed": filing_labels,
            "generated": datetime.now().isoformat(),
            "model": "claude-sonnet-4-6",
        },
        "comparison_steps": [],
    }

    for step in all_comparison_steps:
        merged = step["merged_results"]
        step_data = {
            "current_filing": step["current_label"],
            "compared_against": step["prior_labels"],
            "summary": {
                "total_changes": len(merged),
                "significant": sum(1 for r in merged if r["overall_severity"] == "significant"),
                "moderate": sum(1 for r in merged if r["overall_severity"] == "moderate"),
                "minor": sum(1 for r in merged if r["overall_severity"] == "minor"),
            },
            "pass_summaries": {
                name: {"findings": data["total_findings"], "changes": data["changes_detected"]}
                for name, data in step["pass_results"].items()
            },
            "findings": merged,
        }
        output["comparison_steps"].append(step_data)

    output_path.write_text(json.dumps(output, indent=2), encoding="utf-8")
    print(f"  Full comparison JSON: {output_path}")


def generate_exec_summary_report(
    all_comparison_steps: list,
    deal: DealContext,
    filing_labels: List[str],
    output_path: Path,
    anthropic_key: str = "",
) -> None:
    """Generate a crisp executive summary DOCX: synthesized bullets via Claude."""
    from .comparator import summarize_findings_to_bullets

    print("  Generating executive summary bullets via Claude...")
    bullets = []
    if anthropic_key:
        bullets = summarize_findings_to_bullets(
            all_comparison_steps, deal, filing_labels, anthropic_key
        )

    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(10)
    style.paragraph_format.space_after = Pt(3)
    style.paragraph_format.line_spacing = 1.15

    hdr = doc.add_paragraph()
    r = hdr.add_run(f"{deal.acquirer_company} / {deal.target_company} ({deal.ticker})")
    r.font.size = Pt(13); r.bold = True

    fl = doc.add_paragraph()
    r = fl.add_run(' → '.join(filing_labels))
    r.font.size = Pt(9); r.font.color.rgb = RGBColor(80, 80, 80)

    doc.add_paragraph()

    CATEGORY_COLORS = {
        "Timing":        RGBColor(21, 101, 192),
        "Regulatory":    RGBColor(173, 20, 87),
        "Business/Risk": RGBColor(46, 125, 50),
        "Legal":         RGBColor(60, 60, 60),
        "Other":         RGBColor(60, 60, 60),
    }

    if not bullets:
        p = doc.add_paragraph()
        p.add_run("No material changes identified across all filings.").font.size = Pt(10)
    else:
        headline_items = [b for b in bullets if b.get("category") == "__headline__"]
        body_bullets   = [b for b in bullets if b.get("category") != "__headline__"]

        if headline_items:
            hl_p = doc.add_paragraph()
            hl_p.paragraph_format.space_before = Pt(0)
            hl_p.paragraph_format.space_after = Pt(10)
            hl_r = hl_p.add_run(headline_items[0]["bullet"])
            hl_r.bold = True
            hl_r.font.size = Pt(11)
            hl_r.font.color.rgb = RGBColor(30, 30, 30)

        cat_order = ["Timing", "Regulatory", "Business/Risk", "Legal", "Other"]
        grouped: dict = {}
        for item in body_bullets:
            cat = item.get("category", "Other")
            grouped.setdefault(cat, []).append(item["bullet"])

        for cat in cat_order:
            if cat not in grouped:
                continue
            color = CATEGORY_COLORS.get(cat, RGBColor(60, 60, 60))

            hdr_p = doc.add_paragraph()
            hdr_p.paragraph_format.space_before = Pt(8)
            hdr_p.paragraph_format.space_after = Pt(2)
            hdr_r = hdr_p.add_run(cat.upper())
            hdr_r.bold = True
            hdr_r.font.size = Pt(8)
            hdr_r.font.color.rgb = color

            for bullet in grouped[cat]:
                bp = doc.add_paragraph()
                bp.paragraph_format.left_indent = Inches(0.15)
                bp.paragraph_format.space_before = Pt(2)
                bp.paragraph_format.space_after = Pt(3)

                dot_r = bp.add_run("• ")
                dot_r.font.size = Pt(10)
                dot_r.font.color.rgb = color

                txt_r = bp.add_run(bullet)
                txt_r.font.size = Pt(10)

    doc.save(output_path)
    print(f"  Exec summary saved: {output_path}")
    return bullets
