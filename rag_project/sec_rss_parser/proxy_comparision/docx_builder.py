"""
docx_builder.py — DOCX report generation helpers.
"""

import re

from docx import Document as DocxDocument
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from .report_format import parse_summary_sections, smart_title as _smart_title


# =============================================================================
# Color constants
# =============================================================================

NAVY = RGBColor(0x1F, 0x4E, 0x79)
DARK_GRAY = RGBColor(0x33, 0x33, 0x33)
GREEN = RGBColor(0x00, 0x7A, 0x33)
AMBER = RGBColor(0xBF, 0x8F, 0x00)
LIGHT_GRAY = RGBColor(0x99, 0x99, 0x99)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xCC, 0x00, 0x00)


# =============================================================================
# Helper functions
# =============================================================================

def _set_cell_shading(cell, hex_color: str):
    """Apply background shading to a table cell."""
    shading = OxmlElement('w:shd')
    shading.set(qn('w:val'), 'clear')
    shading.set(qn('w:color'), 'auto')
    shading.set(qn('w:fill'), hex_color)
    cell._tc.get_or_add_tcPr().append(shading)


def _setup_docx_styles(doc: DocxDocument):
    """Configure professional document styles."""
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    style.font.color.rgb = DARK_GRAY
    style.paragraph_format.space_after = Pt(6)

    for level, (size, bold) in {1: (16, True), 2: (13, True), 3: (11, True)}.items():
        h = doc.styles[f'Heading {level}']
        h.font.name = 'Calibri'
        h.font.size = Pt(size)
        h.font.color.rgb = NAVY
        h.font.bold = bold
        h.paragraph_format.space_before = Pt(18 if level == 1 else 12)
        h.paragraph_format.space_after = Pt(6)


def _add_title_page(doc: DocxDocument, ticker: str, target: str,
                     acquirer: str, doc_type_label: str, timestamp: str):
    """Add a professional title page."""
    for _ in range(4):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"{ticker} Merger Filing Analysis")
    run.font.size = Pt(28)
    run.font.color.rgb = NAVY
    run.bold = True

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(doc_type_label)
    run.font.size = Pt(18)
    run.font.color.rgb = DARK_GRAY

    doc.add_paragraph()

    for text in [f"Target: {target}", f"Acquirer: {acquirer}"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.size = Pt(12)
        run.font.color.rgb = DARK_GRAY

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"Generated: {timestamp}")
    run.font.size = Pt(10)
    run.font.color.rgb = LIGHT_GRAY
    run.italic = True

    doc.add_page_break()


def _add_styled_run(paragraph, text: str):
    """Add a run to a paragraph with color coding for [NEW] and arrow markers."""
    if "[NEW]" in text:
        # Split around [NEW] tag
        parts = text.split("[NEW]")
        run = paragraph.add_run(parts[0].strip())
        run.font.size = Pt(11)
        run.font.name = 'Calibri'
        tag_run = paragraph.add_run("  [NEW]")
        tag_run.font.size = Pt(10)
        tag_run.font.color.rgb = GREEN
        tag_run.bold = True
        if len(parts) > 1 and parts[1].strip():
            run2 = paragraph.add_run(" " + parts[1].strip())
            run2.font.size = Pt(11)
            run2.font.name = 'Calibri'
    elif " -> " in text:
        run = paragraph.add_run(text)
        run.font.size = Pt(11)
        run.font.name = 'Calibri'
        run.font.color.rgb = AMBER
    else:
        run = paragraph.add_run(text)
        run.font.size = Pt(11)
        run.font.name = 'Calibri'


# =============================================================================
# Public DOCX builders
# =============================================================================

def create_summary_docx(summary_text: str, output_path: str,
                         ticker: str, target: str, acquirer: str,
                         form_label: str, timestamp: str):
    """Create a professionally formatted summary DOCX from the new section format."""
    doc = DocxDocument()
    _setup_docx_styles(doc)

    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    _add_title_page(doc, ticker, target, acquirer, form_label, timestamp)

    opening, sections = parse_summary_sections(summary_text)

    # Opening paragraph
    if opening:
        p = doc.add_paragraph()
        run = p.add_run(opening)
        run.font.size = Pt(11)
        run.font.name = 'Calibri'
        run.font.color.rgb = DARK_GRAY

    # Each section
    for header, content in sections:
        doc.add_heading(_smart_title(header), level=2)
        if not content.strip():
            continue
        for line in content.split("\n"):
            stripped = line.strip()
            if not stripped:
                continue
            # Bullet items
            if stripped.startswith("- "):
                p = doc.add_paragraph(style='List Bullet')
                _add_styled_run(p, stripped[2:])
            else:
                p = doc.add_paragraph()
                _add_styled_run(p, stripped)

    doc.save(output_path)


def create_changes_docx(change_text: str, output_path: str,
                         ticker: str, target: str, acquirer: str,
                         old_label: str, new_label: str, timestamp: str):
    """Create a professionally formatted changes DOCX from the new section format."""
    doc = DocxDocument()
    _setup_docx_styles(doc)

    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    _add_title_page(doc, ticker, target, acquirer,
                     f"Changes: {old_label} -> {new_label}", timestamp)

    opening, sections = parse_summary_sections(change_text)

    # Opening paragraph
    if opening:
        p = doc.add_paragraph()
        run = p.add_run(opening)
        run.font.size = Pt(11)
        run.font.name = 'Calibri'
        run.font.color.rgb = DARK_GRAY

    # Each section
    for header, content in sections:
        doc.add_heading(_smart_title(header), level=2)
        if not content.strip():
            continue
        for line in content.split("\n"):
            stripped = line.strip()
            if not stripped:
                continue

            # Sub-headers within sections (NEW SENTENCES:, MODIFIED SENTENCES:, etc.)
            sub_match = re.match(r'^(NEW SENTENCES|MODIFIED SENTENCES|REMOVED SENTENCES):?$', stripped)
            if sub_match:
                doc.add_heading(_smart_title(sub_match.group(1)), level=3)
                continue

            # Inserted sentences (+ prefix, green)
            if stripped.startswith("+ "):
                p = doc.add_paragraph(style='List Bullet')
                run = p.add_run(stripped[2:])
                run.font.color.rgb = GREEN
                run.font.size = Pt(9)
                continue

            # Removed sentences (- prefix, red)
            if stripped.startswith("- ") and stripped.startswith("- \""):
                p = doc.add_paragraph(style='List Bullet')
                run = p.add_run(stripped[2:])
                run.font.color.rgb = RED
                run.font.size = Pt(9)
                continue

            # Diff markup lines (OLD:, NEW:)
            if stripped.startswith("OLD: "):
                p = doc.add_paragraph()
                run = p.add_run("OLD: ")
                run.bold = True
                run.font.size = Pt(9)
                run2 = p.add_run(stripped[5:])
                run2.font.size = Pt(9)
                run2.font.color.rgb = LIGHT_GRAY
                continue

            if stripped.startswith("NEW: "):
                p = doc.add_paragraph()
                run = p.add_run("NEW: ")
                run.bold = True
                run.font.size = Pt(9)
                run2 = p.add_run(stripped[5:])
                run2.font.size = Pt(9)
                run2.font.color.rgb = DARK_GRAY
                continue

            # "No changes" / "No material changes" lines
            if ("no changes" in stripped.lower() or "no material changes" in stripped.lower()
                    or "substantially identical" in stripped.lower()
                    or "not found" in stripped.lower()
                    or "sentences unchanged" in stripped.lower()):
                p = doc.add_paragraph()
                run = p.add_run(stripped)
                run.italic = True
                run.font.color.rgb = LIGHT_GRAY
                continue

            # [NEW] items
            if "[NEW]" in stripped:
                p = doc.add_paragraph()
                _add_styled_run(p, stripped)
                continue

            # Arrow changes
            if " -> " in stripped:
                p = doc.add_paragraph()
                _add_styled_run(p, stripped)
                continue

            # Bullet items
            if stripped.startswith("- "):
                p = doc.add_paragraph(style='List Bullet')
                _add_styled_run(p, stripped[2:])
                continue

            # Default: plain paragraph
            p = doc.add_paragraph()
            run = p.add_run(stripped)
            run.font.size = Pt(11)
            run.font.name = 'Calibri'

    doc.save(output_path)
