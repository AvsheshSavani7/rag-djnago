"""
10-K / 10-Q Annual & Quarterly Report Summarizer — Chunked multi-pass via Claude API
Usage: python 10k_summary.py

Handles 10-K, 10-K/A, 10-Q, and 10-Q/A filings using a two-pass architecture:
  Pass 1 (Extraction): Claude Sonnet extracts structured data from each filing section
  Pass 2 (Synthesis):  Claude Sonnet combines section extracts into final L1/L2/L3 output

This approach processes the FULL filing (~60-120 pages) rather than just the first
10,000 words, ensuring complete coverage of financial statements, risk factors, and MD&A.
"""

import anthropic
import time
import re
import json
import sys
import os
import io
from pathlib import Path

try:
    from ._naming import filing_uid, sanitize_date_part, sanitize_filename_part
    from ._deal_context import inject_deal_context
except ImportError:
    from _naming import filing_uid, sanitize_date_part, sanitize_filename_part
    from _deal_context import inject_deal_context

# ──── PASTE YOUR 10-K/10-Q URL HERE ────
FILING_URL = ""
DEAL_CONTEXT = None
# ──── OUTPUT FOLDER (local fallback when not using S3) ────
OUTPUT_DIR = Path(__file__).resolve().parents[1] / "Output Summaries"
# ─────────────────────────────────


try:
    import requests
    from bs4 import BeautifulSoup
    from docx import Document as DocxDocument
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install",
                          "requests", "beautifulsoup4", "python-docx", "-q"])
    import requests
    from bs4 import BeautifulSoup
    from docx import Document as DocxDocument
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

try:
    from ._config import get_anthropic_api_key
except ImportError:
    from _config import get_anthropic_api_key

ANTHROPIC_API_KEY = get_anthropic_api_key()
if not ANTHROPIC_API_KEY and __name__ == "__main__":
    print("❌ ANTHROPIC_API_KEY not found. Set it in .env or Django settings (ANTHROPIC_API_KEY).")
    sys.exit(1)

# ──── MODEL CONFIGURATION ────
CHUNK_MODEL = "claude-sonnet-4-5-20250929"
SYNTHESIS_MODEL = "claude-sonnet-4-5-20250929"
CHUNK_MAX_TOKENS = 1500
SYNTHESIS_MAX_TOKENS = 3000
# ──────────────────────────────

# ──── 10-K ITEM DEFINITIONS ────
ITEMS_10K = {
    "1":  "Business",
    "1A": "Risk Factors",
    "1B": "Unresolved Staff Comments",
    "1C": "Cybersecurity",
    "2":  "Properties",
    "3":  "Legal Proceedings",
    "4":  "Mine Safety Disclosures",
    "5":  "Market for Common Equity",
    "6":  "Reserved",
    "7":  "Management's Discussion and Analysis (MD&A)",
    "7A": "Quantitative and Qualitative Disclosures About Market Risk",
    "8":  "Financial Statements and Supplementary Data",
    "9":  "Changes in and Disagreements With Accountants",
    "9A": "Controls and Procedures",
    "9B": "Other Information",
    "10": "Directors, Executive Officers and Corporate Governance",
    "11": "Executive Compensation",
    "12": "Security Ownership",
    "13": "Certain Relationships and Related Transactions",
    "14": "Principal Accounting Fees and Services",
    "15": "Exhibit and Financial Statement Schedules",
}

# Priority items to extract (skip boilerplate like Properties, Mine Safety)
PRIORITY_ITEMS_10K = ["1", "1A", "3", "7", "7A", "8"]
# 10-Q: Financial Stmts, Risk, MD&A, Market Risk, Controls, Other Info
PRIORITY_ITEMS_10Q = ["1", "1A", "2", "3", "4", "5"]

# Section parsing config
FALLBACK_CHUNK_SIZE = 15000   # words per chunk if section parsing fails
MAX_SECTION_WORDS = 25000     # cap any single section

# Regex to find Item headers in SEC filing text
ITEM_HEADER_RE = re.compile(
    r'(?:^|\n)\s*'
    r'(?:(?:PART\s+(?:I{1,3}|IV)\s*[\.\-—:]*\s*)?'
    r'ITEM\s+(\d+[A-Z]?))\s*[\.\-—:\s]',
    re.IGNORECASE | re.MULTILINE
)


# ──── PROMPTS ────

CHUNK_EXTRACT_PROMPT = """You are an expert financial analyst extracting structured data from a section of an SEC 10-K or 10-Q filing.

You are reading ITEM {item_num} — {item_title}.

Extract ALL quantitative data (dollar amounts, percentages, ratios, share counts, dates) and key qualitative facts from this section. Be thorough — this data will be combined with other sections to produce a final summary.

Respond ONLY in valid JSON (no markdown fences):

{{
  "item": "{item_num}",
  "item_title": "{item_title}",
  "key_facts": ["<the 5-10 most important facts, with exact figures>"],
  "financial_data": {{
    "revenue": "<if mentioned, with period and YoY change>",
    "net_income": "<if mentioned>",
    "eps": "<if mentioned>",
    "margins": "<gross/operating/net margins if mentioned>",
    "operating_income": "<operating income/loss dollar amount with YoY change if available>",
    "operating_cash_flow": "<net cash from operating activities if mentioned>",
    "capital_expenditures": "<capex / purchases of property and equipment if mentioned>",
    "free_cash_flow": "<FCF if stated explicitly, OR calculate as operating cash flow minus capex if both are available>",
    "other_metrics": ["<any other quantitative metrics>"]
  }},
  "risks_identified": ["<specific risks from this section>"],
  "ma_references": ["<any mentions of mergers, acquisitions, divestitures, strategic alternatives>"],
  "forward_looking": ["<guidance, outlook, management expectations>"],
  "notable_changes": ["<anything flagged as new, changed, or materially different from prior period>"]
}}

Rules:
- CRITICAL — FACTS ONLY: Every statement in your summary must be directly traceable to the filing text. Report ONLY what the document says. Do NOT add analysis, assess significance, interpret motives, predict outcomes, evaluate probability, or editorialize. Do NOT state what is "not disclosed" or "not mentioned" — simply omit fields where the filing is silent. If the filing does not say it, do not write it.
  GOOD: "CADE requested revenue data for 2021-2025 across four markets."
  BAD: "The broad scope of information requested indicates potentially detailed competitive analysis ahead."
  GOOD: "The offer expires June 10, 2026."
  BAD: "This tight timeline may create pressure on shareholders to tender quickly."
- PRECISION: Use the filing's exact terminology for legal, regulatory, and financial terms. Do NOT paraphrase in ways that broaden or narrow the stated meaning. GOOD: "All 14 Pennsylvania PUC hearings have concluded." BAD: "Regulatory proceedings concluded in Pennsylvania."
- Extract EXACT numbers — do not round or approximate
- If a data point is not present in this section, use null
- Include both absolute figures and YoY/QoQ changes where stated
- Flag any NEW risk factors or CHANGED disclosures explicitly
- Note any M&A activity, strategic reviews, or activist mentions
- IMPORTANT — Operating income may appear as: "Income (loss) from operations", "Operating income (loss)", "Loss from operations", or similar. Always extract the dollar amount.
- IMPORTANT — Capital expenditures may appear as: "Purchases of property and equipment", "Capital expenditures", "Additions to property", "Purchases of fixed assets", or similar. Always extract from the cash flow statement.
- IMPORTANT — If both operating cash flow and capital expenditures are available, calculate Free Cash Flow = Operating Cash Flow minus Capital Expenditures.

SECTION TEXT:
"""


SYNTHESIS_PROMPT = """You are an expert financial analyst producing a final summary of an SEC {filing_type} filing.

Below are structured extracts from each major section of the filing. Combine them into a single coherent summary with three levels of detail.

The summary should serve BOTH:
1. General financial analysis (earnings, revenue, margins, guidance, balance sheet health)
2. Merger arbitrage relevance (deal status, regulatory conditions, strategic alternatives)

M&A relevance is important but is ONE of many factors — do not force an M&A narrative if the filing is primarily about operating results.

Respond ONLY in valid JSON (no markdown fences):

{{
  "ticker": "<ticker symbol as stated in the filing, or null if not stated>",
  "company": "<company name>",
  "filing_type": "<10-K | 10-K/A | 10-Q | 10-Q/A>",
  "period_end_date": "<fiscal period end, MM/DD/YY>",
  "filing_date": "<filing date, MM/DD/YY>",
  "fiscal_period": "<e.g., FY 2025 or Q3 FY 2025>",

  "L1_headline": "+ <TICKER> – <key takeaway in <=8 words>. | <date>",

  "L2_brief": "<2-3 sentences: key financial result, notable developments, any M&A relevance>",

  "L3_detailed": {{
    "financial_performance": {{
      "revenue": "<total revenue with YoY/QoQ change>",
      "net_income": "<net income with YoY/QoQ change>",
      "eps": "<EPS basic and diluted with YoY/QoQ change>",
      "gross_margin": "<gross margin %>",
      "operating_income": "<operating income with YoY/QoQ change>",
      "free_cash_flow": "<FCF if available>",
      "other_metrics": ["<other key metrics>"]
    }},
    "segment_performance": ["<revenue/growth/margin by segment>"],
    "guidance_and_outlook": "<forward guidance, management outlook>",
    "balance_sheet_highlights": {{
      "cash_and_equivalents": "<cash position>",
      "total_debt": "<total debt>",
      "net_debt": "<net debt>",
      "shareholders_equity": "<total equity>"
    }},
    "risk_factors": ["<top 3-5 material risks, flag NEW or CHANGED>"],
    "ma_and_deal_relevance": {{
      "pending_deals": "<any pending M&A>",
      "deal_impact": "<impact on known deals>",
      "strategic_alternatives": "<strategic review, activism>",
      "acquisition_activity": "<completed/announced acquisitions>"
    }},
    "legal_and_regulatory": ["<material litigation, investigations>"],
    "key_developments": ["<3-5 important non-financial developments>"],
    "capital_allocation": "<dividends, buybacks, capex, debt repayment>"
  }}
}}

Rules:
- CRITICAL — FACTS ONLY: Every statement in your summary must be directly traceable to the filing text. Report ONLY what the document says. Do NOT add analysis, assess significance, interpret motives, predict outcomes, evaluate probability, or editorialize. Do NOT state what is "not disclosed" or "not mentioned" — simply omit fields where the filing is silent. If the filing does not say it, do not write it.
  GOOD: "CADE requested revenue data for 2021-2025 across four markets."
  BAD: "The broad scope of information requested indicates potentially detailed competitive analysis ahead."
  GOOD: "The offer expires June 10, 2026."
  BAD: "This tight timeline may create pressure on shareholders to tender quickly."
- PRECISION: Use the filing's exact terminology for legal, regulatory, and financial terms. Do NOT paraphrase in ways that broaden or narrow the stated meaning. GOOD: "All 14 Pennsylvania PUC hearings have concluded." BAD: "Regulatory proceedings concluded in Pennsylvania."
- L1 format MUST be: + <TICKER> – <takeaway>. | <date>
- Synthesize across sections — do not just concatenate
- Prioritize the MOST material information
- For financial figures, include both absolute values and percentage changes
- If M&A activity is minimal, keep that section brief rather than padding it
- IMPORTANT: For operating_income, look for "Income (loss) from operations" or "Operating income (loss)" in the section extracts. This is a REQUIRED field — never return "Not specified" if any extract contains an operating income figure.
- IMPORTANT: For free_cash_flow, if any extract contains operating_cash_flow and capital_expenditures, CALCULATE FCF = operating cash flow - capex. Only use "Not specified" if neither FCF nor its components appear anywhere in the extracts.
- IMPORTANT: If KNOWN FILING DATE or KNOWN PERIOD END DATE are provided above the extracts, use those exact values for filing_date and period_end_date. Do not guess or leave blank.

SECTION EXTRACTS:
"""


# ──── SECTION PARSING ────

def fetch_filing_full(source: str | list[str]) -> str:
    """Fetch the FULL text of a 10-K/10-Q filing (no word limit truncation).

    Accepts a single URL/path or a list (documents concatenated with separators).
    """
    from .fetch_utils import fetch_text
    if isinstance(source, list):
        parts = []
        for i, s in enumerate(source, 1):
            label = s.split("/")[-1] if "/" in s else s
            print(f"   Fetching document {i}/{len(source)}: {label}")
            parts.append(fetch_text(s, word_limit=0))
        return ("\n\n" + "=" * 60 + "\n\n").join(parts)
    return fetch_text(source, word_limit=0)


def _extract_title(section_text: str, item_num: str) -> str:
    """Extract the section title from the first line of section text.

    E.g., 'Item 1.    Financial Statements (unaudited)' → 'Financial Statements'
    Falls back to ITEMS_10K dictionary lookup.
    """
    # Grab text after "Item N." up to the first newline
    m = re.match(
        r'(?i)\s*(?:PART\s+\S+\s*[\.\-—:]*\s*)?Item\s+\S+\s*[\.\-—:]*\s*(.+)',
        section_text
    )
    if m:
        title_line = m.group(1).strip()
        # Clean: take first meaningful phrase (before next newline or parenthetical)
        title_line = title_line.split('\n')[0].strip()
        # Remove trailing punctuation and parenthetical
        title_line = re.sub(r'\s*\(.*?\)\s*$', '', title_line).strip(' .')
        if len(title_line) > 5 and len(title_line) < 100:
            return title_line

    return ITEMS_10K.get(item_num, f"Item {item_num}")


def _extract_filing_dates(full_text: str, url: str = "") -> dict:
    """Extract filing date and period end date.

    Tries multiple sources:
    1. SGML headers in the raw text (if .txt wrapper was fetched directly)
    2. Visible text patterns like "For the fiscal year ended December 31, 2025"
    3. SGML wrapper .txt file fetched from EDGAR (for filing date on .htm URLs)
    4. URL filename fallback for period end date
    """
    dates = {}
    header = full_text[:10000]

    # Method 1: SGML headers (present if .txt file was fetched)
    m = re.search(r'FILED\s+AS\s+OF\s+DATE[:\s]+(\d{8})', header)
    if m:
        d = m.group(1)
        dates["filing_date"] = f"{d[4:6]}/{d[6:8]}/{d[2:4]}"

    m = re.search(r'CONFORMED\s+PERIOD\s+OF\s+REPORT[:\s]+(\d{8})', header)
    if m:
        d = m.group(1)
        dates["period_end_date"] = f"{d[4:6]}/{d[6:8]}/{d[2:4]}"

    # Method 2: Parse period end date from visible filing text
    if "period_end_date" not in dates:
        m = re.search(
            r'(?:fiscal\s+year|quarterly\s+period|transition\s+period)\s+ended\s+'
            r'(\w+\s+\d{1,2},?\s+\d{4})',
            full_text[:8000], re.IGNORECASE
        )
        if m:
            from datetime import datetime
            try:
                date_str = m.group(1).replace(',', '')
                dt = datetime.strptime(date_str, "%B %d %Y")
                dates["period_end_date"] = dt.strftime("%m/%d/%y")
            except ValueError:
                pass

    # Method 3: Fetch SGML wrapper .txt file for filing date (iXBRL .htm URLs)
    if "filing_date" not in dates and url and '.htm' in url:
        try:
            # .htm URL: .../data/1866368/000186636826000011/cwan-20251231.htm
            # .txt URL: .../data/1866368/000186636826000011/0001866368-26-000011.txt
            m_url = re.search(
                r'(/Archives/edgar/data/\d+/(\d{18})/).*\.htm', url)
            if m_url:
                base_path = m_url.group(1)
                acc_raw = m_url.group(2)  # 000186636826000011
                acc_dashed = f"{acc_raw[:10]}-{acc_raw[10:12]}-{acc_raw[12:]}"
                txt_url = f"https://www.sec.gov{base_path}{acc_dashed}.txt"
                resp = requests.get(txt_url, headers={
                    'User-Agent': 'MergerArbDashboard/1.0 (merger-arb-research@outlook.com)',
                    'Range': 'bytes=0-2000'
                }, timeout=10)
                if resp.status_code in (200, 206):
                    m_filed = re.search(
                        r'FILED\s+AS\s+OF\s+DATE[:\s]+(\d{8})', resp.text)
                    if m_filed:
                        d = m_filed.group(1)
                        dates["filing_date"] = f"{d[4:6]}/{d[6:8]}/{d[2:4]}"
                        print(
                            f"   Filing date from SGML wrapper: {dates['filing_date']}")
                    if "period_end_date" not in dates:
                        m_period = re.search(
                            r'CONFORMED\s+PERIOD\s+OF\s+REPORT[:\s]+(\d{8})', resp.text)
                        if m_period:
                            d = m_period.group(1)
                            dates["period_end_date"] = f"{d[4:6]}/{d[6:8]}/{d[2:4]}"
        except Exception as e:
            print(f"   Could not fetch SGML wrapper for filing date: {e}")

    # Method 4: URL filename fallback for period end (e.g., cwan-20251231.htm)
    if "period_end_date" not in dates and url:
        m = re.search(r'(\d{4})(\d{2})(\d{2})\.htm', url)
        if m:
            dates["period_end_date"] = f"{m.group(2)}/{m.group(3)}/{m.group(1)[2:]}"

    return dates


def _strip_preamble(full_text: str) -> str:
    """Strip XBRL/iXBRL metadata preamble and table of contents.

    Many SEC filings (especially iXBRL) have thousands of words of metadata
    before the actual content. We find the body start by locating the "PART I"
    occurrence that begins the largest block of content (body >> ToC or inline refs).
    """
    # Find all "PART I" occurrences. The \b word boundary already excludes
    # "PART II", "PART III", "PART IV" since "II"/"III"/"IV" are single tokens.
    part1_matches = list(re.finditer(r'(?i)\bPART\s+I\b', full_text))

    if len(part1_matches) >= 2:
        # Pick the PART I that starts the largest content block.
        # Body PART I has thousands of words after it; ToC entries and
        # inline references ("Part I of this Annual Report") are tiny.
        best = max(range(len(part1_matches)), key=lambda i: (
            (part1_matches[i + 1].start() if i + 1 <
             len(part1_matches) else len(full_text))
            - part1_matches[i].start()
        ))
        return full_text[part1_matches[best].start():]
    elif len(part1_matches) == 1:
        return full_text[part1_matches[0].start():]

    # Fallback: find first Item header
    first_item = re.search(r'(?i)\bItem\s+\d', full_text)
    if first_item:
        return full_text[first_item.start():]

    return full_text


def parse_sections(full_text: str) -> list:
    """Parse the filing text into labeled sections based on Item headers.

    Returns a list of dicts:
      [{"item": "1A", "title": "Risk Factors", "text": "...", "word_count": N}, ...]

    Falls back to word-count chunking if fewer than 2 sections found.
    """
    # Strip XBRL preamble and ToC
    body = _strip_preamble(full_text)

    matches = list(ITEM_HEADER_RE.finditer(body))

    if len(matches) < 2:
        return _chunk_by_words(body)

    # Build raw sections between consecutive Item headers
    raw_sections = []
    for i, match in enumerate(matches):
        item_num = match.group(1).upper()
        start = match.start()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(body)
        section_text = body[start:end].strip()
        word_count = len(section_text.split())

        # Extract title from text (e.g., "Item 1. Financial Statements (unaudited)")
        title = _extract_title(section_text, item_num)

        raw_sections.append({
            "item": item_num,
            "title": title,
            "text": section_text,
            "word_count": word_count
        })

    # For duplicate item numbers (ToC vs body, or Part I vs Part II),
    # keep the instance with the most words. ToC entries are tiny (<50 words),
    # body sections are large.
    best_by_item = {}
    for s in raw_sections:
        key = s["item"]
        if key not in best_by_item or s["word_count"] > best_by_item[key]["word_count"]:
            best_by_item[key] = s

    sections = list(best_by_item.values())

    # Cap section size to avoid token overflow
    for s in sections:
        if s["word_count"] > MAX_SECTION_WORDS:
            words = s["text"].split()
            s["text"] = " ".join(words[:MAX_SECTION_WORDS])
            s["word_count"] = MAX_SECTION_WORDS

    # Filter out very small sections (residual ToC fragments)
    sections = [s for s in sections if s["word_count"] >= 30]

    if len(sections) < 2:
        return _chunk_by_words(body)

    # Sort by item number for consistent ordering
    def item_sort_key(s):
        item = s["item"]
        num = re.match(r'(\d+)', item)
        letter = item[len(num.group()):] if num else ""
        return (int(num.group()) if num else 99, letter)

    sections.sort(key=item_sort_key)

    return sections


def select_priority_sections(sections: list, is_10q: bool = False) -> list:
    """Filter to priority sections for API calls.

    Target: 4-6 chunks to keep costs manageable.
    """
    if is_10q:
        priority = set(PRIORITY_ITEMS_10Q)
    else:
        priority = set(PRIORITY_ITEMS_10K)

    selected = [s for s in sections if s["item"] in priority]

    # If we found fewer than 2 priority sections, take all sections
    if len(selected) < 2:
        selected = sections

    return selected


def _chunk_by_words(text: str) -> list:
    """Fallback: split text into chunks of FALLBACK_CHUNK_SIZE words."""
    words = text.split()
    chunks = []
    for i in range(0, len(words), FALLBACK_CHUNK_SIZE):
        chunk_words = words[i:i + FALLBACK_CHUNK_SIZE]
        chunk_num = (i // FALLBACK_CHUNK_SIZE) + 1
        chunks.append({
            "item": f"CHUNK_{chunk_num}",
            "title": f"Section chunk {chunk_num}",
            "text": " ".join(chunk_words),
            "word_count": len(chunk_words)
        })
    return chunks


def detect_filing_type(text: str) -> str:
    """Detect whether this is a 10-K, 10-K/A, 10-Q, or 10-Q/A from the text."""
    first_2000 = " ".join(text.split()[:2000]).upper()
    if "10-K/A" in first_2000 or "10-KA" in first_2000:
        return "10-K/A"
    if "10-Q/A" in first_2000 or "10-QA" in first_2000:
        return "10-Q/A"
    if "10-Q" in first_2000 or "QUARTERLY REPORT" in first_2000:
        return "10-Q"
    return "10-K"


def detect_ticker(url: str, text: str) -> str:
    """Detect the ticker symbol for the filing company.

    Strategy (in order):
    1. Perplexity API lookup from company name (most reliable)
    2. Extract from URL filename as fallback
    """
    # Extract company name from filing text for Perplexity lookup
    company_name = ""
    body = _strip_preamble(text)
    # Common patterns: "HOLOGIC, INC." or "APPLE INC" on its own line near top
    m = re.search(
        r'(?:^|\n)\s*([A-Z][A-Z\s&,\.]{2,50}(?:INC|CORP|LLC|LP|LTD|CO|GROUP|HOLDINGS)\.?)\s*(?:\n|$)', body[:3000])
    if m:
        company_name = m.group(1).strip().rstrip('.')

    # Method 1: Perplexity API lookup (fast, cheap, reliable)
    if company_name:
        try:
            pplx_key = os.getenv("PERPLEXITY_API_KEY")
            if pplx_key:
                resp = requests.post(
                    'https://api.perplexity.ai/chat/completions',
                    headers={
                        'Authorization': f'Bearer {pplx_key}',
                        'Content-Type': 'application/json'
                    },
                    json={
                        'model': 'sonar',
                        'messages': [{
                            'role': 'user',
                            'content': f'What is the stock ticker symbol for {company_name}? Reply with ONLY the ticker symbol, nothing else.'
                        }],
                        'max_tokens': 20,
                        'temperature': 0
                    },
                    timeout=10
                )
                if resp.status_code == 200:
                    raw = resp.json()["choices"][0]["message"]["content"]
                    # Clean: extract just the ticker from response (may have markdown/citations)
                    ticker_match = re.search(r'[A-Z]{1,6}', raw)
                    if ticker_match:
                        ticker = ticker_match.group(0)
                        print(
                            f"   Ticker via Perplexity: {ticker} (company: {company_name})")
                        return ticker
        except Exception as e:
            print(f"   Perplexity lookup failed: {e}")

    # Method 2: Extract from URL filename (fallback)
    # Pattern: /holx-20251227.htm → HOLX
    m = re.search(r'/([a-zA-Z]{2,6})[-_]\d{6,8}\.htm', url)
    if m:
        ticker = m.group(1).upper()
        print(f"   Ticker from URL: {ticker}")
        return ticker

    return ""


# ──── CORE PROCESSING ────

def extract_section(client, section: dict, retries: int = 2) -> dict:
    """Pass 1: Use Sonnet to extract structured data from a single section.

    Retries on JSON parse failure for reliability.
    """
    prompt = CHUNK_EXTRACT_PROMPT.format(
        item_num=section["item"],
        item_title=section["title"]
    )

    for attempt in range(retries + 1):
        msg = client.messages.create(
            model=CHUNK_MODEL,
            max_tokens=CHUNK_MAX_TOKENS,
            temperature=0,
            messages=[{
                "role": "user",
                "content": prompt + "\n\n" + section["text"]
            }]
        )

        raw = msg.content[0].text.strip()
        raw = re.sub(r"^```json\s*", "", raw)
        raw = re.sub(r"\s*```$", "", raw)

        try:
            return json.loads(raw)
        except json.JSONDecodeError:
            if attempt < retries:
                print(
                    f"     Retry {attempt + 1}: JSON parse failed, retrying...")
                time.sleep(1)
            else:
                return {"item": section["item"], "raw_extract": raw, "parse_error": True}


def synthesize_extracts(client, extracts: list, filing_type: str, ticker_hint: str = "", filing_dates: dict = None, deal_context: dict | None = None) -> dict:
    """Pass 2: Combine section extracts into final L1/L2/L3 summary."""
    extracts_text = json.dumps(extracts, indent=2)

    prompt = inject_deal_context(SYNTHESIS_PROMPT.format(
        filing_type=filing_type), deal_context)

    # Inject known metadata so the model doesn't have to guess
    metadata_lines = []
    if ticker_hint:
        metadata_lines.append(f"KNOWN TICKER: {ticker_hint}")
    if filing_dates:
        if "filing_date" in filing_dates:
            metadata_lines.append(
                f"KNOWN FILING DATE: {filing_dates['filing_date']}")
        if "period_end_date" in filing_dates:
            metadata_lines.append(
                f"KNOWN PERIOD END DATE: {filing_dates['period_end_date']}")
    if metadata_lines:
        extracts_text = "\n".join(metadata_lines) + "\n\n" + extracts_text

    msg = client.messages.create(
        model=SYNTHESIS_MODEL,
        max_tokens=SYNTHESIS_MAX_TOKENS,
        temperature=0,
        messages=[{
            "role": "user",
            "content": prompt + "\n\n" + extracts_text
        }]
    )

    raw = msg.content[0].text.strip()
    raw = re.sub(r"^```json\s*", "", raw)
    raw = re.sub(r"\s*```$", "", raw)

    return json.loads(raw)


def summarize(text: str, model: str = None, deal_context: dict | None = None, filing_url=None) -> dict:
    """Full multi-pass summarization pipeline.

    1. Detect filing type (10-K vs 10-Q)
    2. Parse sections from full text
    3. Select priority sections
    4. Extract data from each section (Pass 1 — Sonnet)
    5. Synthesize into final summary (Pass 2 — Sonnet)

    The 'model' parameter is accepted for interface compatibility with the
    router but is not used — models are controlled by CHUNK_MODEL and
    SYNTHESIS_MODEL constants.
    """
    if not ANTHROPIC_API_KEY:
        raise ValueError(
            "ANTHROPIC_API_KEY not set. Set it in .env or Django settings (ANTHROPIC_API_KEY).")
    client = anthropic.Anthropic(api_key=ANTHROPIC_API_KEY)

    # Step 0: Detect ticker and filing dates from header/URL
    filing_url_ref = (
        filing_url[0] if isinstance(filing_url, list) else filing_url
    ) if filing_url is not None else (
        FILING_URL[0] if isinstance(FILING_URL, list) else FILING_URL
    )
    ticker_hint = detect_ticker(filing_url_ref, text)
    filing_dates = _extract_filing_dates(text, filing_url_ref)
    if filing_dates:
        print(f"   Dates from header: {filing_dates}")

    # Step 1: Detect filing type
    filing_type = detect_filing_type(text)
    is_10q = "Q" in filing_type
    print(f"   Detected filing type: {filing_type}")

    # Step 2: Parse sections
    all_sections = parse_sections(text)
    print(f"   Parsed {len(all_sections)} sections from filing")
    for s in all_sections:
        print(
            f"     Item {s['item']}: {s['title']} ({s['word_count']:,} words)")

    # Step 3: Select priority sections
    sections = select_priority_sections(all_sections, is_10q=is_10q)
    total_words = sum(s["word_count"] for s in sections)
    print(
        f"   Selected {len(sections)} priority sections ({total_words:,} words total)")

    # Step 4: Extract from each section (Pass 1)
    extracts = []
    for i, section in enumerate(sections):
        print(f"   [{i+1}/{len(sections)}] Extracting Item {section['item']}: "
              f"{section['title']} ({section['word_count']:,} words)...")
        extract = extract_section(client, section)
        extracts.append(extract)
        if i < len(sections) - 1:
            time.sleep(0.5)  # Rate limit courtesy

    print(f"   Pass 1 complete: {len(extracts)} section extracts")

    # Step 5: Synthesize (Pass 2)
    print(f"   Synthesizing final summary via {SYNTHESIS_MODEL}...")
    result = synthesize_extracts(
        client, extracts, filing_type, ticker_hint=ticker_hint,
        filing_dates=filing_dates, deal_context=deal_context)

    return result


# ──── OUTPUT ────

def print_summary(s: dict):
    """Pretty-print the multi-level summary."""
    print("\n" + "=" * 70)
    filing_type = s.get("filing_type", "10-K/10-Q")
    print(f"  {filing_type} SUMMARY")
    print("=" * 70)

    print(
        f"\n   Company:    {s.get('company', 'N/A')} ({s.get('ticker', 'N/A')})")
    print(f"   Type:       {filing_type}")
    print(f"   Period:     {s.get('fiscal_period', 'N/A')}")
    print(f"   Period End: {s.get('period_end_date', 'N/A')}")
    print(f"   Filed:      {s.get('filing_date', 'N/A')}")

    print(f"\n   L1 | HEADLINE")
    print(f"   {s['L1_headline']}")

    print(f"\n   L2 | BRIEF")
    print(f"   {s['L2_brief']}")

    d = s["L3_detailed"]
    fp = d.get("financial_performance", {})
    print(f"\n   L3 | DETAILED")

    print(f"\n   -- FINANCIAL PERFORMANCE --")
    print(f"   Revenue:      {fp.get('revenue', 'N/A')}")
    print(f"   Net Income:   {fp.get('net_income', 'N/A')}")
    print(f"   EPS:          {fp.get('eps', 'N/A')}")
    print(f"   Gross Margin: {fp.get('gross_margin', 'N/A')}")
    print(f"   Op Income:    {fp.get('operating_income', 'N/A')}")
    print(f"   FCF:          {fp.get('free_cash_flow', 'N/A')}")
    if fp.get("other_metrics"):
        for m in fp["other_metrics"]:
            print(f"     - {m}")

    if d.get("segment_performance"):
        print(f"\n   -- SEGMENT PERFORMANCE --")
        for seg in d["segment_performance"]:
            print(f"     - {seg}")

    if d.get("guidance_and_outlook"):
        print(f"\n   -- GUIDANCE & OUTLOOK --")
        print(f"   {d['guidance_and_outlook']}")

    bs = d.get("balance_sheet_highlights", {})
    print(f"\n   -- BALANCE SHEET --")
    print(f"   Cash:     {bs.get('cash_and_equivalents', 'N/A')}")
    print(f"   Debt:     {bs.get('total_debt', 'N/A')}")
    print(f"   Net Debt: {bs.get('net_debt', 'N/A')}")
    print(f"   Equity:   {bs.get('shareholders_equity', 'N/A')}")

    if d.get("risk_factors"):
        print(f"\n   -- RISK FACTORS --")
        for r in d["risk_factors"]:
            print(f"     - {r}")

    ma = d.get("ma_and_deal_relevance", {})
    print(f"\n   -- M&A & DEAL RELEVANCE --")
    print(f"   Pending Deals:  {ma.get('pending_deals', 'N/A')}")
    print(f"   Deal Impact:    {ma.get('deal_impact', 'N/A')}")
    print(f"   Strategic Alts: {ma.get('strategic_alternatives', 'N/A')}")
    print(f"   Acq Activity:   {ma.get('acquisition_activity', 'N/A')}")

    if d.get("legal_and_regulatory"):
        print(f"\n   -- LEGAL & REGULATORY --")
        for item in d["legal_and_regulatory"]:
            print(f"     - {item}")

    if d.get("key_developments"):
        print(f"\n   -- KEY DEVELOPMENTS --")
        for k in d["key_developments"]:
            print(f"     - {k}")

    if d.get("capital_allocation"):
        print(f"\n   -- CAPITAL ALLOCATION --")
        print(f"   {d['capital_allocation']}")

    print("=" * 70)


def export_docx(s: dict, s3_key_suffix: str):
    """Build summary as Word doc, upload to S3 (summary_docx/), return (s3_path, s3_url)."""
    from .s3_utils import upload_docx_bytes

    ticker = s.get("ticker", "UNKNOWN")
    filing_type = s.get("filing_type", "10-K")
    date = s.get("filing_date", "")

    doc = DocxDocument()

    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)

    title = doc.add_heading(
        f"{filing_type} Summary: {s.get('company', ticker)}", level=0)
    title.runs[0].font.size = Pt(20)

    # Metadata
    meta = doc.add_paragraph()
    meta.add_run("Company: ").bold = True
    meta.add_run(f"{s.get('company', 'N/A')} ({ticker})")
    meta.add_run("    Period: ").bold = True
    meta.add_run(s.get("fiscal_period", "N/A"))

    meta2 = doc.add_paragraph()
    meta2.add_run("Period End: ").bold = True
    meta2.add_run(s.get("period_end_date", "N/A"))
    meta2.add_run("    Filing Date: ").bold = True
    meta2.add_run(date)

    # L1
    doc.add_heading("L1 — Headline", level=1)
    p = doc.add_paragraph()
    run = p.add_run(s["L1_headline"])
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0, 51, 102)

    # L2
    doc.add_heading("L2 — Brief", level=1)
    doc.add_paragraph(s["L2_brief"])

    # L3
    doc.add_heading("L3 — Detailed", level=1)
    d = s["L3_detailed"]
    fp = d.get("financial_performance", {})

    doc.add_heading("Financial Performance", level=2)
    fp_p = doc.add_paragraph()
    fp_p.add_run("Revenue: ").bold = True
    fp_p.add_run(str(fp.get("revenue", "N/A")) + "\n")
    fp_p.add_run("Net Income: ").bold = True
    fp_p.add_run(str(fp.get("net_income", "N/A")) + "\n")
    fp_p.add_run("EPS: ").bold = True
    fp_p.add_run(str(fp.get("eps", "N/A")) + "\n")
    fp_p.add_run("Gross Margin: ").bold = True
    fp_p.add_run(str(fp.get("gross_margin", "N/A")) + "\n")
    fp_p.add_run("Operating Income: ").bold = True
    fp_p.add_run(str(fp.get("operating_income", "N/A")) + "\n")
    fp_p.add_run("Free Cash Flow: ").bold = True
    fp_p.add_run(str(fp.get("free_cash_flow", "N/A")))

    if fp.get("other_metrics"):
        doc.add_heading("Other Financial Metrics", level=3)
        for m in fp["other_metrics"]:
            doc.add_paragraph(str(m), style="List Bullet")

    if d.get("segment_performance"):
        doc.add_heading("Segment Performance", level=2)
        for seg in d["segment_performance"]:
            doc.add_paragraph(str(seg), style="List Bullet")

    if d.get("guidance_and_outlook"):
        doc.add_heading("Guidance & Outlook", level=2)
        doc.add_paragraph(str(d["guidance_and_outlook"]))

    bs = d.get("balance_sheet_highlights", {})
    doc.add_heading("Balance Sheet", level=2)
    bs_p = doc.add_paragraph()
    bs_p.add_run("Cash & Equivalents: ").bold = True
    bs_p.add_run(str(bs.get("cash_and_equivalents", "N/A")) + "\n")
    bs_p.add_run("Total Debt: ").bold = True
    bs_p.add_run(str(bs.get("total_debt", "N/A")) + "\n")
    bs_p.add_run("Net Debt: ").bold = True
    bs_p.add_run(str(bs.get("net_debt", "N/A")) + "\n")
    bs_p.add_run("Shareholders' Equity: ").bold = True
    bs_p.add_run(str(bs.get("shareholders_equity", "N/A")))

    if d.get("risk_factors"):
        doc.add_heading("Risk Factors", level=2)
        for r in d["risk_factors"]:
            doc.add_paragraph(str(r), style="List Bullet")

    ma = d.get("ma_and_deal_relevance", {})
    doc.add_heading("M&A & Deal Relevance", level=2)
    ma_p = doc.add_paragraph()
    ma_p.add_run("Pending Deals: ").bold = True
    ma_p.add_run(str(ma.get("pending_deals", "N/A")) + "\n")
    ma_p.add_run("Deal Impact: ").bold = True
    ma_p.add_run(str(ma.get("deal_impact", "N/A")) + "\n")
    ma_p.add_run("Strategic Alternatives: ").bold = True
    ma_p.add_run(str(ma.get("strategic_alternatives", "N/A")) + "\n")
    ma_p.add_run("Acquisition Activity: ").bold = True
    ma_p.add_run(str(ma.get("acquisition_activity", "N/A")))

    if d.get("legal_and_regulatory"):
        doc.add_heading("Legal & Regulatory", level=2)
        for item in d["legal_and_regulatory"]:
            doc.add_paragraph(str(item), style="List Bullet")

    if d.get("key_developments"):
        doc.add_heading("Key Developments", level=2)
        for k in d["key_developments"]:
            doc.add_paragraph(str(k), style="List Bullet")

    if d.get("capital_allocation"):
        doc.add_heading("Capital Allocation", level=2)
        doc.add_paragraph(str(d["capital_allocation"]))

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    path, url = upload_docx_bytes(buf.read(), s3_key_suffix)
    return path, url


# ──── MAIN ────

def main(filing_url=None, deal_context: dict | None = None):
    ctx = deal_context if deal_context is not None else DEAL_CONTEXT
    source = filing_url if filing_url is not None else FILING_URL
    if isinstance(source, list):
        print(f"Fetching 10-K/10-Q from {len(source)} document(s)")
        for u in source:
            print(f"   • {u.split('/')[-1]}")
    else:
        print(f"Fetching 10-K/10-Q from: {source}")

    text = fetch_filing_full(source)
    total_words = len(text.split())
    print(f"Extracted {total_words:,} words of text (full document)")

    print("Starting chunked multi-pass summarization...")
    result = summarize(text, deal_context=ctx, filing_url=source)
    from ._ticker_context import apply_known_tickers
    result = apply_known_tickers(result, ctx)

    print_summary(result)

    uid = filing_uid(source)
    from .s3_utils import upload_json

    s3_json_path, s3_json_url = upload_json(result, f"10k_summary_{uid}.json")
    print(f"\nJSON uploaded to S3: {s3_json_url}")

    safe_ticker = sanitize_filename_part(result.get("ticker"))
    safe_type = sanitize_filename_part(result.get("filing_type"), "10-K")
    safe_date = sanitize_date_part(result.get("filing_date"))
    docx_suffix = f"{safe_type}_Summary_{safe_ticker}_{safe_date}_{uid}.docx"
    s3_docx_path, s3_docx_url = export_docx(result, docx_suffix)
    print(f"DOCX uploaded to S3: {s3_docx_url}")

    result["s3_docx_path"] = s3_docx_path
    result["s3_docx_url"] = s3_docx_url
    result["s3_json_path"] = s3_json_path
    result["s3_json_url"] = s3_json_url
    return result


if __name__ == "__main__":
    main()
